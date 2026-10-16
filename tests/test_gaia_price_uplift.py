import tempfile
import unittest
import os
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from flask import Flask
from playwright.sync_api import sync_playwright
from app.gaia_listing.api.routes import create_gaia_blueprint
from app.gaia_listing.application.price_uplift import (
    PRICE_RULE_NAMES, PRICE_RULE_RATES, PriceUpliftError, prepare_price_uplift,
    rule_name_matches_shop_term,
)
from app.gaia_listing.domain.models import GaiaBatch
from app.gaia_listing.infrastructure.batch_repository import GaiaBatchRepository
from app.gaia_listing.infrastructure.price_uplift_template import (
    inspect_price_uplift_template, write_price_uplift_template,
)
from integrations.gaia_browser.price_rule_upload import (
    _import_product_button, _price_file_button, _response_rejection, _response_summary,
)
from integrations.gaia_browser.price_rule_probe import locate_price_rule


SHEET_XML = b'''<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A1:B1"/><sheetData><row r="1">
<c r="A1" t="inlineStr"><is><t>\xe6\x8a\x96\xe9\x9f\xb3\xe5\x95\x86\xe5\x93\x81ID\xef\xbc\x88\xe5\xbf\x85\xe5\xa1\xab\xef\xbc\x89</t></is></c>
<c r="B1" t="inlineStr"><is><t>\xe5\xaf\xbc\xe5\x85\xa5\xe5\xa4\xb1\xe8\xb4\xa5\xe5\x8e\x9f\xe5\x9b\xa0</t></is></c>
</row></sheetData></worksheet>'''


def make_template(path):
    with ZipFile(path, 'w') as archive:
        archive.writestr('xl/worksheets/sheet1.xml', SHEET_XML)


class PriceUpliftTests(unittest.TestCase):
    def test_browser_rule_selection_uses_id_and_term_and_excludes_live_room(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.set_content('''<table><tbody>
                  <tr><td>已改名的店铺【51165770】 -12期</td><td>启用</td><td>浮动比例：4.3%</td></tr>
                  <tr><td>已改名的店铺直播间【51165770】 -12期</td><td>启用</td><td>浮动比例：4.3%</td></tr>
                  <tr><td>已改名的店铺【51165770】 -24期</td><td>启用</td><td>浮动比例：9%</td></tr>
                  <tr><td>其他店铺【7970406】 -12期</td><td>启用</td><td>浮动比例：4.3%</td></tr>
                </tbody></table>''')
                row, cells = locate_price_rule(page, '51165770', 12)
                self.assertEqual(cells[0], '已改名的店铺【51165770】 -12期')
                self.assertIn('已改名的店铺', row.inner_text())
                page.locator('tbody').evaluate('node => node.insertAdjacentHTML("beforeend", '
                    '"<tr><td>另一条货架规则【51165770】-12期</td><td>启用</td><td>浮动比例：4.3%</td></tr>")')
                with self.assertRaisesRegex(PriceUpliftError, '命中 2 条'):
                    locate_price_rule(page, '51165770', 12)
            finally:
                browser.close()

    def test_rule_identity_uses_bracketed_shop_id_and_exact_term_not_shop_name(self):
        self.assertTrue(rule_name_matches_shop_term(
            '返航新生二手手机旗舰店【208501840】 -12期', '208501840', 12))
        self.assertTrue(rule_name_matches_shop_term(
            '已改名的店铺【208501840】—24期', '208501840', 24))
        self.assertFalse(rule_name_matches_shop_term(
            '已改名的店铺【208501840】—24期', '208501840', 12))
        self.assertFalse(rule_name_matches_shop_term(
            '返航新生二手手机旗舰店【208501841】—12期', '208501840', 12))
        self.assertFalse(rule_name_matches_shop_term(
            '返航新生二手手机旗舰店【208501840】【123】—12期', '208501840', 12))
        self.assertFalse(rule_name_matches_shop_term(
            '返航新生二手手机旗舰店 208501840—12期', '208501840', 12))
        self.assertFalse(rule_name_matches_shop_term(
            '爱回收严选二手优品直播间【51165770】-12期', '51165770', 12))

    def test_import_button_waits_for_async_edit_form(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.set_content('<div id="actions"></div>')
                page.evaluate('''() => setTimeout(() => {
                    document.querySelector('#actions').innerHTML =
                        '<button type="button"><span>导入商品ID</span></button>';
                }, 250)''')
                button = _import_product_button(page, timeout_seconds=2)
                self.assertEqual('BUTTON', button.evaluate('(element) => element.tagName'))
                page.evaluate('''() => document.querySelector('#actions').insertAdjacentHTML(
                    'beforeend', '<button type="button">导入商品ID</button>')''')
                with self.assertRaisesRegex(PriceUpliftError, '2 个可见'):
                    _import_product_button(page, timeout_seconds=0.2)
            finally:
                browser.close()

    def test_submit_response_diagnostics_strip_query_and_expose_business_error(self):
        class Response:
            url = 'https://gaiya.jd.com/api/rule/import?access_token=secret'
            status = 200

            @staticmethod
            def json():
                return {'code': 40004, 'msg': '商品 ID 不符合规则'}

        summary = _response_summary(Response())
        self.assertEqual('/api/rule/import', summary['path'])
        self.assertNotIn('secret', str(summary))
        self.assertEqual('商品 ID 不符合规则', _response_rejection([summary]))
        class GaiaResponse(Response):
            url = ('https://gaiya.jd.com/api?apiName='
                   'pop.erp.live.rule.import.goods&access_token=secret')

        gaia_summary = _response_summary(GaiaResponse())
        self.assertEqual('pop.erp.live.rule.import.goods', gaia_summary['api_name'])
        self.assertNotIn('secret', str(gaia_summary))
        self.assertEqual('商品 ID 不符合规则', _response_rejection([gaia_summary]))
        self.assertEqual('', _response_rejection([
            {'path': '/api/telemetry/report', 'status': 500, 'msg': 'unrelated'},
        ]))

    def test_nested_upload_wrapper_has_one_native_file_button(self):
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.set_content('''<div role="dialog"><span class="ant-upload" role="button">
                    <button type="button" class="ant-btn"><span>选择文件</span></button>
                    </span></div>''')
                dialog = page.get_by_role('dialog')
                self.assertEqual(2, dialog.get_by_role('button', name='选择文件').count())
                self.assertEqual('BUTTON', _price_file_button(dialog).evaluate('(element) => element.tagName'))
            finally:
                browser.close()

    def test_verified_shops_have_separate_twelve_and_twenty_four_month_rules(self):
        for shop_id, rules in PRICE_RULE_NAMES.items():
            with self.subTest(shop_id=shop_id):
                self.assertEqual({12, 24}, set(rules))
                self.assertNotEqual(rules[12], rules[24])
                self.assertIn('12期', rules[12])
                self.assertIn('24期', rules[24])
                self.assertEqual({12, 24}, set(PRICE_RULE_RATES[shop_id]))
        self.assertEqual('拍拍严选手机优品【168976032】-12期', PRICE_RULE_NAMES['168976032'][12])
        self.assertEqual('拍拍严选二手手机【5576543】—24期', PRICE_RULE_NAMES['5576543'][24])
        self.assertEqual('返航新生二手手机旗舰店【208501840】-12期', PRICE_RULE_NAMES['208501840'][12])
        self.assertEqual({12: '4.3%', 24: '9%'}, PRICE_RULE_RATES['7970406'])
        self.assertTrue(rule_name_matches_shop_term('直播间-运营店【7970406】-6期', '7970406', 6, 1))
        self.assertFalse(rule_name_matches_shop_term('直播间-运营店【7970406】-6期', '7970406', 6, 2))
        self.assertFalse(rule_name_matches_shop_term('运营店【7970406】-6期', '7970406', 6, 1))

    def test_only_reconciled_same_shop_batch_is_ready(self):
        batch = {'id': 'abc123', 'shop_id': '7970406', 'scene': 2, 'state': 'gaia_accepted',
                 'summary': {'generated_rows': 18, 'expected_product_count': 2},
                 'doudian_draft_ids': ['3843713525176467501', '3843710140591243616']}
        with self.assertRaisesRegex(PriceUpliftError, '逐行'):
            prepare_price_uplift(batch, shop_id='7970406', installments=12)
        batch['state'] = 'doudian_drafts_matched'
        with self.assertRaisesRegex(PriceUpliftError, '店铺'):
            prepare_price_uplift(batch, shop_id='156450040', installments=12)
        self.assertEqual(6, prepare_price_uplift(batch, shop_id='7970406', installments=6)['installments'])
        batch['scene'] = 1
        self.assertEqual(12, prepare_price_uplift(batch, shop_id='7970406', installments=12)['installments'])
        batch['scene'] = 2
        prepared = prepare_price_uplift(batch, shop_id='7970406', installments=24)
        self.assertEqual(2, prepared['count'])
        self.assertEqual('店铺 ID 7970406 · 24 期（实际规则名以盖亚页面为准）', prepared['rule_name'])
        batch['summary']['expected_product_count'] = 1
        with self.assertRaisesRegex(PriceUpliftError, '预计商品数'):
            prepare_price_uplift(batch, shop_id='7970406', installments=24)
        batch['summary']['expected_product_count'] = 2
        batch['doudian_draft_ids'][1] = batch['doudian_draft_ids'][0]
        with self.assertRaisesRegex(PriceUpliftError, '重复'):
            prepare_price_uplift(batch, shop_id='7970406', installments=12)
        batch['shop_id'] = '999999999'
        with self.assertRaisesRegex(PriceUpliftError, '授权店铺'):
            prepare_price_uplift(batch, shop_id='999999999', installments=12)

    def test_template_copies_a_column_and_preserves_b_header(self):
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / 'template.xlsx'
            output = Path(root) / 'output.xlsx'
            make_template(source)
            inspect_price_uplift_template(str(source))
            write_price_uplift_template(str(source), str(output), ['3843713525176467501'])
            with ZipFile(source) as archive:
                self.assertNotIn(b'A2', archive.read('xl/worksheets/sheet1.xml'))
            with ZipFile(output) as archive:
                xml = archive.read('xl/worksheets/sheet1.xml')
                self.assertIn(b'A2', xml)
                self.assertIn(b'3843713525176467501', xml)
                self.assertIn(b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">', xml)
                self.assertNotIn(b'<s:worksheet', xml)

    def test_generated_workbook_preserves_official_xml_and_is_excel_readable(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / 'template.xlsx'
            output = Path(root) / 'output.xlsx'
            book = Workbook()
            sheet = book.active
            sheet['A1'] = '抖音商品ID（必填）'
            sheet['B1'] = '导入失败原因'
            book.save(source)
            write_price_uplift_template(str(source), str(output), ['3843713525176467501'])
            with ZipFile(output) as archive:
                xml = archive.read('xl/worksheets/sheet1.xml')
                self.assertIn(b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">', xml)
                self.assertNotIn(b'<s:worksheet', xml)
            self.assertEqual('3843713525176467501', load_workbook(output).active['A2'].value)

    def test_preview_blocks_accepted_batch_without_ids(self):
        with tempfile.TemporaryDirectory() as root:
            app = Flask('price_uplift_test')
            app.register_blueprint(create_gaia_blueprint(root, ['7970406']))
            repository = GaiaBatchRepository(str(Path(root) / 'panel_data/gaia_listing/batches'))
            repository.save(GaiaBatch('abc123', '7970406', 2, 1, '/source', 'hash',
                                      state='gaia_accepted', summary={'generated_rows': 1}))
            response = app.test_client().post('/api/gaia/batches/abc123/price-uplift/preview',
                                              json={'shop_id': '7970406', 'installments': 12})
            self.assertEqual(409, response.status_code)
            self.assertIn('逐行', response.get_json()['message'])

    def test_latest_official_download_is_managed_without_modifying_source(self):
        with tempfile.TemporaryDirectory() as root:
            downloads = Path(root) / 'downloads'
            downloads.mkdir()
            source = downloads / '直播商品价格浮动规则数据模版20261009140756.xlsx'
            make_template(source)
            original = source.read_bytes()
            with patch.dict(os.environ, {'GAIA_DOWNLOADS_DIR': str(downloads)}):
                app = Flask('price_uplift_download_test')
                app.register_blueprint(create_gaia_blueprint(root, ['7970406']))
            client = app.test_client()
            response = client.get('/api/gaia/price-uplift/template/current')
            self.assertEqual(200, response.status_code, response.get_json())
            self.assertEqual('downloads', response.get_json()['source'])
            self.assertEqual(original, source.read_bytes())
            self.assertEqual('managed', client.get('/api/gaia/price-uplift/template/current').get_json()['source'])


if __name__ == '__main__':
    unittest.main()
