import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook, load_workbook

import panel_server
from app.gaia_listing.api.routes import create_gaia_blueprint
from app.gaia_listing.domain.template_contract import GAIA_TEMPLATE_HEADERS
from app.gaia_listing.domain.models import GaiaBatch
from app.gaia_listing.infrastructure.batch_repository import GaiaBatchRepository


class GaiaListingApiTests(unittest.TestCase):
    def test_managed_template_upload_validation_versioning_and_legacy_migration(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            config = root_path / 'config'
            config.mkdir()
            project = Path(__file__).resolve().parents[1]
            for name in ('gaia_excel_profiles.json', 'gaia_quality_mapping.json'):
                (config / name).write_bytes((project / 'config' / name).read_bytes())
            legacy_dir = root_path / 'panel_data/gaia_listing/imports/templates'
            legacy_dir.mkdir(parents=True)
            valid = Workbook()
            valid.active.append(list(GAIA_TEMPLATE_HEADERS))
            legacy = legacy_dir / 'old-template.xlsx'
            valid.save(legacy)
            legacy_bytes = legacy.read_bytes()
            app = Flask('gaia_managed_template_test')
            app.register_blueprint(create_gaia_blueprint(root))
            client = app.test_client()

            migrated = client.get('/api/gaia/templates/current').get_json()['template']
            self.assertEqual('legacy_managed', migrated['source_kind'])
            self.assertEqual('old-template.xlsx', migrated['name'])
            self.assertEqual(legacy_bytes, legacy.read_bytes())
            self.assertEqual(migrated, client.get('/api/gaia/templates/current').get_json()['template'])

            wrong = Workbook()
            wrong.active.append(['wrong', 'headers'])
            wrong_stream = io.BytesIO()
            wrong.save(wrong_stream)
            denied = client.post('/api/gaia/templates', data={
                'file': (io.BytesIO(wrong_stream.getvalue()), 'wrong.xlsx'),
            }, content_type='multipart/form-data')
            self.assertEqual(400, denied.status_code)
            self.assertEqual(migrated, client.get('/api/gaia/templates/current').get_json()['template'])

            uploaded = client.post('/api/gaia/templates', data={
                'file': (io.BytesIO(legacy_bytes), 'new-template.xlsx'),
            }, content_type='multipart/form-data')
            self.assertEqual(201, uploaded.status_code, uploaded.get_json())
            current = uploaded.get_json()['template']
            self.assertNotEqual(migrated['id'], current['id'])
            self.assertEqual(current, client.get('/api/gaia/templates/current').get_json()['template'])
            self.assertEqual(legacy_bytes, Path(migrated['path']).read_bytes())
            self.assertEqual(legacy_bytes, Path(current['path']).read_bytes())

            Path(current['path']).write_bytes(b'tampered')
            corrupted = client.get('/api/gaia/templates/current')
            self.assertEqual(409, corrupted.status_code)
            self.assertIn('哈希不一致', corrupted.get_json()['message'])

    def test_gaia_quality_codes_preview_excel_and_upload_guard(self):
        """Use shipped Gaia config, not synthetic Doudian property option IDs."""
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            config_dir = root_path / 'config'
            config_dir.mkdir()
            project = Path(__file__).resolve().parents[1]
            for name in ('gaia_excel_profiles.json', 'gaia_quality_mapping.json'):
                (config_dir / name).write_bytes((project / 'config' / name).read_bytes())
            downloads = root_path / 'downloads'
            downloads.mkdir()
            template = downloads / 'template.xlsx'
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            workbook.save(template)
            original = template.read_bytes()
            with patch.dict(os.environ, {'GAIA_DOWNLOADS_DIR': str(downloads)}):
                app = Flask('gaia_quality_codes')
                app.register_blueprint(create_gaia_blueprint(root, ['5576543']))
            client = app.test_client()
            mapping = client.post('/api/gaia/model-mappings', data={
                'file': (io.BytesIO('品牌,型号,标品ID,品类\n苹果,iPhone 16,123456,手机\n'.encode()), 'mapping.csv'),
            }, content_type='multipart/form-data').get_json()['mapping']
            repo = GaiaBatchRepository(str(root_path / 'panel_data/gaia_listing/batches'))
            for label, code, wrong_id in [('99新', '99', '18972'), ('95新', '95', '18954'),
                                          ('9新', '90', '9'), ('准新机', '100', '0')]:
                with self.subTest(label=label):
                    payload = {'shop_id': '5576543', 'source_path': str(template),
                               'mapping_version_id': mapping['id'], 'quality_label': label,
                               'selected_sku_ids': ['123456'], 'scene': 2, 'create_count': 1}
                    response = client.post('/api/gaia/batches/preview', json=payload)
                    self.assertEqual(200, response.status_code, response.get_json())
                    preview = response.get_json()['preview']
                    self.assertEqual(code, preview['quality_id'])
                    self.assertEqual(code, preview['rows'][0]['values'][4])
                    self.assertEqual('a', preview['rows'][0]['values'][7])
                    payload.update(expected_source_sha256=preview['source_sha256'], expected_quality_id=wrong_id)
                    self.assertEqual(400, client.post('/api/gaia/batches/generate', json=payload).status_code)
                    payload['expected_quality_id'] = code
                    generated = client.post('/api/gaia/batches/generate', json=payload)
                    self.assertEqual(201, generated.status_code, generated.get_json())
                    batch = generated.get_json()['batch']
                    self.assertEqual(code, batch['quality_id'])
                    download = client.get(generated.get_json()['download_url'])
                    self.assertEqual(200, download.status_code)
                    book = load_workbook(io.BytesIO(download.data), read_only=True, data_only=True)
                    self.assertEqual(['2', 'A', '1', '123456', code, None, 'D且E且G', 'a'],
                                     [book.active.cell(2, column).value for column in range(1, 9)])
                    book.close()
                    download.close()
                    url = '/api/gaia/batches/%s/handoff-preflight?shop_id=5576543' % batch['id']
                    self.assertEqual(200, client.get(url).status_code)
                    # A persisted batch made with the old code cannot be sent.
                    repo.save(GaiaBatch(**dict(batch, quality_id=wrong_id)))
                    blocked = client.get(url)
                    self.assertEqual(409, blocked.status_code)
                    self.assertIn('成色 ID 在生成后发生变化', '；'.join(blocked.get_json()['preflight']['errors']))
            multi = {'shop_id': '5576543', 'source_path': str(template),
                     'mapping_version_id': mapping['id'], 'quality_labels': ['99新', '95新', '9新', '准新机'],
                     'selected_sku_ids': ['123456'], 'scene': 2, 'create_count': 1}
            preview_response = client.post('/api/gaia/batches/preview', json=multi)
            self.assertEqual(200, preview_response.status_code, preview_response.get_json())
            preview = preview_response.get_json()['preview']
            self.assertEqual(4, preview['row_count'])
            self.assertEqual(['99', '95', '90', '100'], preview['quality_ids'])
            self.assertEqual(['99', '95', '90', '100'], [row['values'][4] for row in preview['rows']])
            multi.update(expected_source_sha256=preview['source_sha256'],
                         expected_quality_ids=preview['quality_ids'])
            generated = client.post('/api/gaia/batches/generate', json=multi)
            self.assertEqual(201, generated.status_code, generated.get_json())
            batch = generated.get_json()['batch']
            self.assertEqual(['99新', '95新', '9新', '准新机'], batch['quality_labels'])
            self.assertEqual(['99', '95', '90', '100'], batch['quality_ids'])
            self.assertEqual(1, batch['summary']['expected_product_count'])
            download = client.get(generated.get_json()['download_url'])
            book = load_workbook(io.BytesIO(download.data), read_only=True, data_only=True)
            self.assertEqual(['99', '95', '90', '100'], [book.active.cell(row, 5).value for row in (2, 3, 4, 5)])
            book.close()
            download.close()
            self.assertEqual(200, client.get('/api/gaia/batches/%s/handoff-preflight?shop_id=5576543'
                                             % batch['id']).status_code)
            self.assertEqual(original, template.read_bytes(), 'source template must stay unchanged')

    def test_stop_endpoint_is_scoped_idempotent_and_refuses_submitted_work(self):
        with tempfile.TemporaryDirectory() as root:
            app = Flask('gaia_stop_test')
            app.register_blueprint(create_gaia_blueprint(root, ['7970406', '5576543']))
            client = app.test_client()
            repo = GaiaBatchRepository(str(Path(root) / 'panel_data/gaia_listing/batches'))
            batch = GaiaBatch('abc123', '7970406', 2, 1, '/source', 'hash',
                              state='validation_passed', output_hash='hash')
            repo.save(batch)
            url = '/api/gaia/batches/abc123/stop'
            self.assertEqual(400, client.post(url, json=['bad']).status_code)
            self.assertEqual(400, client.post(url, json={'shop_id': '5576543', 'expected_output_hash': 'hash'}).status_code)
            self.assertEqual(400, client.post(url, json={'shop_id': '7970406', 'expected_output_hash': 'wrong'}).status_code)
            payload = {'shop_id': '7970406', 'expected_output_hash': 'hash'}
            stopped = client.post(url, json=payload)
            self.assertEqual(200, stopped.status_code)
            self.assertEqual('stopped', stopped.get_json()['batch']['gaia_receipt']['phase'])
            self.assertEqual(stopped.get_json(), client.post(url, json=payload).get_json())
            batch.state = 'gaia_submitted'
            batch.gaia_receipt = {'phase': 'submitting', 'submit_intent': True}
            repo.save(batch)
            self.assertEqual(409, client.post(url, json=payload).status_code)
            self.assertEqual('gaia_submitted', repo.get('abc123')['state'])

    def setUp(self):
        panel_server.app.config.update(TESTING=True)
        self.client = panel_server.app.test_client()

    def test_capabilities_fail_closed(self):
        with patch.dict(os.environ, {"GAIA_LISTING_V1_ENABLED": "0"}):
            response = self.client.get("/api/gaia/capabilities")
        self.assertEqual(200, response.status_code)
        data = response.get_json()
        self.assertFalse(data["features"]["gaia_submit"])
        self.assertTrue(data["safety"]["gaia_submit_requires_confirmation"])
        self.assertTrue(data["features"]["model_mapping_upload"])
        self.assertTrue(data["features"]["template_generation"])
        self.assertTrue(data["features"]["gaia_handoff_preflight"])
        self.assertEqual("explicit_mainland_only_v2", data["contracts"]["mapping_policy"])

    def test_upload_valid_csv_and_restore_latest_version(self):
        with tempfile.TemporaryDirectory() as root:
            app = Flask(__name__)
            app.register_blueprint(create_gaia_blueprint(root, ["208501840"]))
            client = app.test_client()
            source = "品牌,型号,标品ID,品类\n荣耀,400 Pro,123456,手机\n".encode("utf-8")
            response = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO(source), "models.csv"),
            }, content_type="multipart/form-data")
            self.assertEqual(201, response.status_code)
            mapping = response.get_json()["mapping"]
            self.assertEqual(1, mapping["record_count"])
            self.assertEqual({"手机": 1}, mapping["category_counts"])
            self.assertEqual(mapping, client.get("/api/gaia/model-mappings/latest").get_json()["mapping"])
            restarted = Flask("restarted_gaia_listing")
            restarted.register_blueprint(create_gaia_blueprint(root, ["208501840"]))
            self.assertEqual(
                mapping,
                restarted.test_client().get("/api/gaia/model-mappings/latest").get_json()["mapping"],
            )

            invalid = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO("品牌,型号,标品ID\n,400 Pro,123456\n".encode("utf-8")), "bad.csv"),
            }, content_type="multipart/form-data")
            self.assertEqual(400, invalid.status_code)
            self.assertIn("缺少品牌", invalid.get_json()["message"])
            self.assertEqual(mapping, client.get("/api/gaia/model-mappings/latest").get_json()["mapping"])

            stored = Path(root) / "panel_data" / "gaia_listing" / "imports" / "model_mappings" / (mapping["id"] + ".csv")
            stored.write_bytes(b"changed")
            changed = client.get("/api/gaia/model-mappings/latest")
            self.assertEqual(409, changed.status_code)
            self.assertIn("哈希不一致", changed.get_json()["message"])

    def test_upload_rejects_duplicate_sku_and_xls(self):
        with tempfile.TemporaryDirectory() as root:
            app = Flask(__name__)
            app.register_blueprint(create_gaia_blueprint(root))
            client = app.test_client()
            duplicate = "品牌,型号,标品ID\n荣耀,400 Pro,123456\n荣耀,400,123456\n".encode("utf-8")
            response = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO(duplicate), "models.csv"),
            }, content_type="multipart/form-data")
            self.assertEqual(400, response.status_code)
            self.assertIn("同一标品ID对应多个型号", response.get_json()["message"])
            response = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO(b"not-an-xlsx"), "models.xls"),
            }, content_type="multipart/form-data")
            self.assertEqual(400, response.status_code)
            response = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO(b"not-an-xlsx"), "models.xlsx"),
            }, content_type="multipart/form-data")
            self.assertEqual(400, response.status_code)
            self.assertIsNone(client.get("/api/gaia/model-mappings/latest").get_json()["mapping"])

    def test_publish_template_is_not_a_model_mapping_and_old_unfiltered_copy_is_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            app = Flask("gaia_wrong_mapping")
            app.register_blueprint(create_gaia_blueprint(root))
            client = app.test_client()
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            stream = io.BytesIO()
            workbook.save(stream)
            stream.seek(0)
            response = client.post("/api/gaia/model-mappings", data={
                "file": (stream, "双维成色发品模板.xlsx"),
            }, content_type="multipart/form-data")
            self.assertEqual(400, response.status_code)
            self.assertIn("不是型号映射表", response.get_json()["message"])
            self.assertIsNone(client.get("/api/gaia/model-mappings/latest").get_json()["mapping"])

            old = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO("品牌,型号,标品ID\n苹果,iPhone 17,123\n".encode()), "old.csv"),
            }, content_type="multipart/form-data").get_json()["mapping"]
            path = Path(root) / "panel_data/gaia_listing/imports/model_mappings" / (old["id"] + ".csv")
            response = client.post("/api/gaia/model-mappings/import-download", json={"path": str(path)})
            self.assertEqual(400, response.status_code)
            self.assertIn("国行筛选", response.get_json()["message"])

    def test_imports_recognized_mapping_from_downloads_without_changing_source(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            downloads = root_path / "downloads"
            downloads.mkdir()
            config = root_path / "config"
            config.mkdir()
            profile_source = Path(__file__).resolve().parents[1] / "config" / "gaia_excel_profiles.json"
            (config / "gaia_excel_profiles.json").write_bytes(profile_source.read_bytes())
            (config / "gaia_quality_mapping.json").write_text('{"qualities": []}', encoding="utf-8")
            source = downloads / "抖音直播货架货品价格规则数据导出(货品维度)-生产20260914140744.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["类目", "品牌", "型号", "货品ID", "拍拍标品skuId", "标品名称", "一维成色"])
            sheet.append(["手机", "苹果", "iPhone 17", "goods-1", "sku-1", "苹果 iPhone 17 国行", "99新"])
            workbook.save(source)
            original = source.read_bytes()
            with patch.dict(os.environ, {"GAIA_DOWNLOADS_DIR": str(downloads)}):
                app = Flask("gaia_download_mapping_test")
                app.register_blueprint(create_gaia_blueprint(root))
            client = app.test_client()
            response = client.post("/api/gaia/model-mappings/import-download", json={"path": str(source)})
            self.assertEqual(201, response.status_code, response.get_json())
            mapping = response.get_json()["mapping"]
            self.assertEqual(1, mapping["record_count"])
            self.assertTrue(mapping["mainland_filter_applied"])
            self.assertEqual("explicit_mainland_only_v2", mapping["mapping_policy"])
            self.assertEqual(0, mapping["excluded_non_mainland_rows"])
            self.assertEqual("downloads", mapping["source_kind"])
            self.assertEqual(original, source.read_bytes())

            stored = root_path / "panel_data/gaia_listing/imports/model_mappings" / (mapping["id"] + ".xlsx")
            from app.gaia_listing.infrastructure.excel_inspector import discover_excel_files

            def discover_without_downloads(directory):
                if str(directory) == str(downloads):
                    raise PermissionError("Operation not permitted: Downloads")
                return discover_excel_files(directory)

            with patch("app.gaia_listing.api.routes.discover_excel_files", side_effect=discover_without_downloads):
                listing = client.get("/api/gaia/excel/latest").get_json()
            self.assertEqual("downloads", listing["warnings"][0]["source"])
            self.assertEqual(str(stored.resolve()), listing["files"][0]["path"])
            reused = client.post("/api/gaia/model-mappings/import-download", json={"path": str(stored)})
            self.assertEqual(200, reused.status_code, reused.get_json())
            self.assertTrue(reused.get_json()["reused"])
            self.assertEqual(mapping, reused.get_json()["mapping"])
            self.assertEqual(1, len(list(stored.parent.glob("*.xlsx"))))

            outside = root_path / "outside.xlsx"
            outside.write_bytes(original)
            denied = client.post("/api/gaia/model-mappings/import-download", json={"path": str(outside)})
            self.assertEqual(400, denied.status_code)
            self.assertIn("允许的导入目录", denied.get_json()["message"])

            unknown = stored.parent / "unregistered.xlsx"
            unknown.write_bytes(original)
            denied = client.post("/api/gaia/model-mappings/import-download", json={"path": str(unknown)})
            self.assertEqual(400, denied.status_code)

            resolved = client.post("/api/gaia/model-mappings/resolve", json={
                "mapping_version_id": mapping["id"], "brand": "Apple", "model": "Apple iPhone 17",
            })
            self.assertEqual(200, resolved.status_code, resolved.get_json())
            self.assertEqual("sku-1", resolved.get_json()["resolution"]["candidates"][0]["standard_sku_id"])

            stored.write_bytes(b"tampered")
            denied = client.post("/api/gaia/model-mappings/import-download", json={"path": str(stored)})
            self.assertEqual(400, denied.status_code)
            self.assertIn("哈希不一致", denied.get_json()["message"])
            stored.write_bytes(original)
            index = stored.with_suffix(".index.json")
            index.write_bytes(b"[]")
            denied = client.post("/api/gaia/model-mappings/import-download", json={"path": str(stored)})
            self.assertEqual(400, denied.status_code)
            self.assertIn("索引哈希不一致", denied.get_json()["message"])

    def test_operator_models_batch_resolve_only_when_every_sku_is_unique(self):
        with tempfile.TemporaryDirectory() as root:
            app = Flask("gaia_model_batch_resolution")
            app.register_blueprint(create_gaia_blueprint(root))
            client = app.test_client()
            source = (
                "品牌,型号,标品ID,品类\n"
                "苹果,iPhone 17 Pro,sku-apple,手机\n"
                "荣耀,Magic8 Pro,sku-honor,手机\n"
                "甲牌,Common X,sku-a,手机\n"
                "乙牌,Common X,sku-b,手机\n"
            ).encode("utf-8")
            upload = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO(source), "models.csv"),
            }, content_type="multipart/form-data")
            self.assertEqual(201, upload.status_code)
            mapping_id = upload.get_json()["mapping"]["id"]

            matched = client.post("/api/gaia/model-mappings/resolve-batch", json={
                "mapping_version_id": mapping_id,
                "models": [
                    {"brand": "Apple", "model": "Apple iPhone 17 Pro"},
                    {"brand": "", "model": "Magic8 Pro"},
                ],
            })
            self.assertEqual(200, matched.status_code, matched.get_json())
            self.assertEqual(["sku-apple", "sku-honor"],
                             matched.get_json()["selected_sku_ids"])

            ambiguous = client.post("/api/gaia/model-mappings/resolve-batch", json={
                "mapping_version_id": mapping_id,
                "models": [{"brand": "", "model": "Common X"}],
            })
            self.assertEqual(409, ambiguous.status_code)
            data = ambiguous.get_json()
            self.assertFalse(data["all_unique"])
            self.assertEqual(2, data["results"][0]["candidate_count"])
            self.assertEqual([], data["selected_sku_ids"])

    def test_excel_discovery_falls_back_to_managed_template_when_downloads_unavailable(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            missing_downloads = root_path / "not-readable-downloads"
            config = root_path / "config"
            config.mkdir()
            source_config = Path(__file__).resolve().parents[1] / "config"
            (config / "gaia_excel_profiles.json").write_bytes(
                (source_config / "gaia_excel_profiles.json").read_bytes())
            (config / "gaia_quality_mapping.json").write_text('{"qualities": []}', encoding="utf-8")
            managed = root_path / "panel_data" / "gaia_listing" / "imports" / "templates"
            managed.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append([*GAIA_TEMPLATE_HEADERS, "说明"])
            workbook.save(managed / "抖音双维成色发品模版.xlsx")
            with patch.dict(os.environ, {"GAIA_DOWNLOADS_DIR": str(missing_downloads)}):
                app = Flask("gaia_managed_template_test")
                app.register_blueprint(create_gaia_blueprint(root))
            response = app.test_client().get("/api/gaia/excel/latest")
            self.assertEqual(200, response.status_code)
            files = response.get_json()["files"]
            self.assertEqual(1, len(files))
            self.assertEqual("managed_template", files[0]["source_kind"])
            self.assertEqual("gaia_publish_template", files[0]["profile_matches"][0]["id"])

    def test_preview_generation_download_and_quality_fail_closed(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            downloads = root_path / "downloads"
            downloads.mkdir()
            config = root_path / "config"
            config.mkdir()
            profile_source = Path(__file__).resolve().parents[1] / "config" / "gaia_excel_profiles.json"
            (config / "gaia_excel_profiles.json").write_bytes(profile_source.read_bytes())
            quality_path = config / "gaia_quality_mapping.json"
            quality_path.write_text(json.dumps({"qualities": [
                {"label": "99新", "dy_quality_id": ""},
                {"label": "95新", "dy_quality_id": ""},
            ]}), encoding="utf-8")
            template = downloads / "盖亚双维成色发品模板.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append([*GAIA_TEMPLATE_HEADERS, "盖亚原始说明"])
            workbook.save(template)
            original = template.read_bytes()
            with patch.dict(os.environ, {"GAIA_DOWNLOADS_DIR": str(downloads)}):
                app = Flask("gaia_b2_test")
                app.register_blueprint(create_gaia_blueprint(root, ["208501840", "156450040"]))
            client = app.test_client()
            source = "品牌,型号,标品ID,品类\n荣耀,400 Pro,123456,手机\n苹果,iPad Air,987654,平板\n".encode()
            upload = client.post("/api/gaia/model-mappings", data={
                "file": (io.BytesIO(source), "models.csv"),
            }, content_type="multipart/form-data")
            self.assertEqual(201, upload.status_code)
            mapping = upload.get_json()["mapping"]
            payload = {"shop_id": "208501840", "scene": 2, "create_count": 1,
                       "source_path": str(template), "mapping_version_id": mapping["id"],
                       "quality_label": "99新", "selected_sku_ids": ["123456", "987654"]}
            blocked = client.post("/api/gaia/batches/preview", json=payload)
            self.assertEqual(400, blocked.status_code)
            self.assertIn("dyQualityId", blocked.get_json()["message"])
            self.assertFalse((root_path / "panel_data" / "gaia_listing" / "outputs").exists())
            quality_path.write_text(json.dumps({"qualities": [
                {"label": "99新", "dy_quality_id": "quality-99"},
                {"label": "95新", "dy_quality_id": ""},
            ]}), encoding="utf-8")
            statuses = client.get("/api/gaia/qualities/status").get_json()["qualities"]
            self.assertTrue(statuses["99新"]["configured"])
            self.assertFalse(statuses["95新"]["configured"])
            preview_response = client.post("/api/gaia/batches/preview", json=payload)
            self.assertEqual(200, preview_response.status_code)
            preview = preview_response.get_json()["preview"]
            self.assertEqual(2, preview["row_count"])
            self.assertEqual("D且E且G", preview["rows"][0]["values"][6])
            self.assertEqual("", preview["rows"][1]["values"][6])
            self.assertEqual(400, client.post("/api/gaia/batches/generate", json=payload).status_code)
            payload.update(expected_source_sha256=preview["source_sha256"],
                           expected_quality_id=preview["quality_id"])
            stale = dict(payload, expected_source_sha256="0" * 64)
            self.assertEqual(400, client.post("/api/gaia/batches/generate", json=stale).status_code)
            generated = client.post("/api/gaia/batches/generate", json=payload)
            self.assertEqual(201, generated.status_code, generated.get_json())
            self.assertEqual(original, template.read_bytes())
            batch = generated.get_json()["batch"]
            self.assertEqual("validation_passed", batch["state"])
            self.assertEqual(2, batch["summary"]["generated_rows"])
            download = client.get(generated.get_json()["download_url"])
            self.assertEqual(200, download.status_code, download.get_json())
            result = load_workbook(io.BytesIO(download.data), read_only=True, data_only=True)
            self.assertEqual("盖亚原始说明", result.active["I1"].value)
            self.assertEqual("quality-99", result.active["E2"].value)
            self.assertEqual("D且E且G", result.active["G2"].value)
            self.assertEqual("a", result.active["H2"].value)
            self.assertEqual("a", result.active["H3"].value)
            self.assertIsNone(result.active["G3"].value)
            result.close()
            download.close()
            preflight_url = "/api/gaia/batches/%s/handoff-preflight?shop_id=208501840" % batch["id"]
            handoff = client.get(preflight_url)
            self.assertEqual(200, handoff.status_code)
            manifest = handoff.get_json()["preflight"]
            self.assertTrue(manifest["ready"])
            self.assertTrue(manifest["requires_human_confirmation"])
            self.assertFalse(manifest["gaia_submitted"])
            self.assertEqual(409, client.get(preflight_url.replace("208501840", "156450040")).status_code)
            quality_path.write_text(json.dumps({"qualities": [
                {"label": "99新", "dy_quality_id": "changed-id"},
            ]}), encoding="utf-8")
            self.assertEqual(409, client.get(preflight_url).status_code)
            quality_path.write_text(json.dumps({"qualities": [
                {"label": "99新", "dy_quality_id": "quality-99"},
            ]}), encoding="utf-8")
            template.write_bytes(b"changed-template")
            self.assertEqual(409, client.get(preflight_url).status_code)
            template.write_bytes(original)
            mapping_path = root_path / "panel_data" / "gaia_listing" / "imports" / "model_mappings" / (mapping["id"] + ".csv")
            mapping_original = mapping_path.read_bytes()
            mapping_path.write_bytes(b"changed-mapping")
            self.assertEqual(409, client.get(preflight_url).status_code)
            mapping_path.write_bytes(mapping_original)
            self.assertEqual(200, client.get(preflight_url).status_code)
            Path(batch["output_path"]).write_bytes(b"tampered")
            self.assertEqual(409, client.get(generated.get_json()["download_url"]).status_code)
            self.assertEqual(409, client.get(preflight_url).status_code)


if __name__ == "__main__":
    unittest.main()
