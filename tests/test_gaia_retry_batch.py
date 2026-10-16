import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook

from app.gaia_listing.application.create_batch import sha256_file
from app.gaia_listing.application.retry_batch import create_remaining_batch
from app.gaia_listing.domain.mapping_contract import ModelMappingRecord
from app.gaia_listing.domain.models import GaiaBatch
from app.gaia_listing.domain.template_contract import GAIA_TEMPLATE_HEADERS, GaiaTemplateRow
from app.gaia_listing.infrastructure.batch_repository import GaiaBatchRepository
from app.gaia_listing.infrastructure.excel_inspector import load_profiles
from app.gaia_listing.infrastructure.template_writer import write_generated_template


class GaiaRetryBatchTests(unittest.TestCase):
    def test_only_rejected_variant_is_removed_from_new_immutable_batch(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            source = root / "抖音双维成色发品模板.xlsx"
            book = Workbook()
            book.active.title = "Sheet1"
            book.active.append(list(GAIA_TEMPLATE_HEADERS))
            book.save(source)
            outputs = root / "outputs"
            old_output = outputs / "abc123" / "gaia_upload.xlsx"
            rows = [GaiaTemplateRow("2", sku, grade, "手机",
                                   aggregate_mark="A" if sku == "111" else "B",
                                   quality_link="0")
                    for sku in ("111", "222") for grade in ("99", "95")]
            write_generated_template(template_path=str(source), output_path=str(old_output), rows=rows)
            original_hash = sha256_file(old_output)
            repository = GaiaBatchRepository(root / "batches")
            old = GaiaBatch(
                "abc123", "5576543", 2, 1, str(source), sha256_file(source),
                state="gaia_rejected", output_path=str(old_output.resolve()),
                output_hash=original_hash, mapping_version_id="v1", mapping_hash="maphash",
                quality_label="99新、95新", quality_id="99,95",
                quality_labels=["99新", "95新"], quality_ids=["99", "95"],
                summary={"generated_rows": 4, "blocking_errors": 0},
            ).to_dict()
            models = [ModelMappingRecord("小米", "18 Pro", "111", "手机"),
                      ModelMappingRecord("苹果", "18 Pro", "222", "手机")]
            mapping = Mock(load_records=Mock(return_value=({"id": "v1", "sha256": "maphash"}, models)))
            profiles = load_profiles(str(Path(__file__).resolve().parents[1]
                                         / "config" / "gaia_excel_profiles.json"))
            child = create_remaining_batch(
                previous=old, rejected_pair=("111", "95"), repository=repository,
                outputs_dir=outputs, mapping_repository=mapping, profiles=profiles,
            )
            self.assertEqual("validation_passed", child["state"])
            self.assertEqual(3, child["summary"]["generated_rows"])
            self.assertEqual(2, child["summary"]["expected_product_count"])
            self.assertEqual(original_hash, sha256_file(old_output), "old submitted bytes must not change")
            workbook = load_workbook(child["output_path"], read_only=True, data_only=True)
            try:
                self.assertEqual([("111", "99"), ("222", "99"), ("222", "95")],
                                 [(str(row[3]), str(row[4])) for row in
                                  workbook.active.iter_rows(min_row=2, max_col=8, values_only=True)])
                self.assertEqual([("A", "0"), ("B", "0"), ("B", "0")],
                                 [(str(row[1]), str(row[2])) for row in
                                  workbook.active.iter_rows(min_row=2, max_col=8, values_only=True)])
            finally:
                workbook.close()
            with self.assertRaisesRegex(ValueError, "未在原批次中唯一出现"):
                create_remaining_batch(previous=old, rejected_pair=("999", "95"),
                                       repository=repository, outputs_dir=outputs,
                                       mapping_repository=mapping, profiles=profiles)


if __name__ == "__main__":
    unittest.main()
