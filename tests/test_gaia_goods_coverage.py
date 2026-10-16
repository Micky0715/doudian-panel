import io
import json
import tempfile
import unittest
from pathlib import Path

from flask import Flask
from openpyxl import Workbook, load_workbook

from app.gaia_listing.api.routes import create_gaia_blueprint
from app.gaia_listing.application.generate_batch import (
    prepare_gaia_template, generate_gaia_batch, split_model_parts,
)
from app.gaia_listing.domain.template_contract import GAIA_TEMPLATE_HEADERS
from app.gaia_listing.infrastructure.excel_inspector import load_profiles
from app.gaia_listing.infrastructure.model_mapping_import import _records_from_rows
from app.gaia_listing.infrastructure.model_mapping_repository import ModelMappingRepository
from app.gaia_listing.infrastructure.batch_repository import GaiaBatchRepository
from app.gaia_listing.infrastructure.quality_mapping_repository import QualityMappingRepository


class GaiaGoodsCoverageTests(unittest.TestCase):
    def test_each_model_gets_its_own_file_even_below_500_rows(self):
        from types import SimpleNamespace
        models = [SimpleNamespace(brand='小米', model='18 Pro')] * 300
        models += [SimpleNamespace(brand='小米', model='18 Ultra')] * 199
        models += [SimpleNamespace(brand='小米', model='18 Pro Max')]
        self.assertEqual(((0, 300), (300, 499), (499, 500)), split_model_parts(models, 499))
        short = [SimpleNamespace(brand='iQOO', model='12 Pro')] * 36
        short += [SimpleNamespace(brand='iQOO', model='12')] * 54
        self.assertEqual(((0, 36), (36, 90)), split_model_parts(short, 499))

    def test_500_data_rows_are_blocked_before_any_gaia_upload(self):
        with tempfile.TemporaryDirectory() as root:
            project = Path(__file__).resolve().parents[1]
            mapping_path = Path(root) / 'mapping.csv'
            mapping_path.write_text(
                '类目,品牌,型号,货品ID,拍拍标品skuId,标品名称,双维成色\n' +
                ''.join('手机,小米,小米 18 Pro,%s,%s,小米 18 Pro 国行,99S\n' % (i, i)
                        for i in range(1, 501)), encoding='utf-8',
            )
            mapping_repo = ModelMappingRepository(Path(root) / 'versions')
            mapping = mapping_repo.save_path(mapping_path)
            template = Path(root) / 'template.xlsx'
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            workbook.save(template)
            with self.assertRaisesRegex(ValueError, '单型号有 500 行'):
                prepare_gaia_template(
                    source_path=template, mapping_version_id=mapping['id'],
                    quality_label=['99新'], scene=2,
                    selected_sku_ids=[str(i) for i in range(1, 501)],
                    profiles=load_profiles(str(project / 'config/gaia_excel_profiles.json')),
                    mapping_repository=mapping_repo,
                    quality_repository=QualityMappingRepository(
                        str(project / 'config/gaia_quality_mapping.json')),
                )

    def test_three_models_auto_split_into_three_verified_files(self):
        with tempfile.TemporaryDirectory() as root:
            project = Path(__file__).resolve().parents[1]
            mapping_path = Path(root) / 'mapping.csv'
            models = ('小米 18 Pro', '小米 18 Pro Max', '小米 18 Ultra')
            mapping_path.write_text(
                '类目,品牌,型号,货品ID,拍拍标品skuId,标品名称,双维成色\n' +
                ''.join('手机,小米,%s,%s,%s,%s 国行,99S\n' % (
                    model, number, number, model)
                    for group, model in enumerate(models)
                    for number in range(group * 200 + 1, (group + 1) * 200 + 1)),
                encoding='utf-8',
            )
            mapping_repo = ModelMappingRepository(Path(root) / 'versions')
            mapping = mapping_repo.save_path(mapping_path)
            template = Path(root) / 'template.xlsx'
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            workbook.save(template)
            profiles = load_profiles(str(project / 'config/gaia_excel_profiles.json'))
            prepared = prepare_gaia_template(
                source_path=template, mapping_version_id=mapping['id'],
                quality_label=['99新'], scene=2,
                selected_sku_ids=[str(i) for i in range(1, 601)], profiles=profiles,
                mapping_repository=mapping_repo,
                quality_repository=QualityMappingRepository(
                    str(project / 'config/gaia_quality_mapping.json')),
            )
            self.assertEqual(((0, 200), (200, 400), (400, 600)), prepared.parts)
            self.assertEqual(3, prepared.preview()['file_count'])
            repository = GaiaBatchRepository(Path(root) / 'batches')
            first = generate_gaia_batch(
                repository=repository, outputs_dir=Path(root) / 'outputs',
                shop_id='5576543', scene=2, create_count=1,
                source_path=template, profiles=profiles, prepared=prepared,
            )
            ids = first.summary['split_batch_ids']
            self.assertEqual(3, len(ids))
            self.assertEqual([200, 200, 200], [repository.get(value)['summary']['generated_rows'] for value in ids])
            for batch_id, count in zip(ids, (200, 200, 200)):
                book = load_workbook(repository.get(batch_id)['output_path'], read_only=True)
                try:
                    self.assertEqual(count + 1, book.active.max_row)
                finally:
                    book.close()

    def test_live_gaia_count_splits_one_model_and_generation_uses_all_chunks(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            config = root_path / 'config'
            config.mkdir()
            project = Path(__file__).resolve().parents[1]
            for name in ('gaia_excel_profiles.json', 'gaia_quality_mapping.json'):
                (config / name).write_bytes((project / 'config' / name).read_bytes())
            template = root_path / 'panel_data/gaia_listing/imports/template.xlsx'
            template.parent.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            workbook.save(template)
            calls = []

            def probe(models):
                calls.append(models)
                return {'source': 'test', 'checked_at': '2026-10-16T12:00:00+08:00',
                        'models': [{**models[0], 'total_count': 1007}]}

            app = Flask('gaia_live_goods_split')
            app.register_blueprint(create_gaia_blueprint(
                root, ['5576543'], goods_count_probe=probe))
            client = app.test_client()
            csv = ('类目,品牌,型号,货品ID,拍拍标品skuId,标品名称,一维成色,双维成色\n' +
                   ''.join('手机,华为,华为 Mate X6,%s,sku-%s,华为 Mate X6 国行,,99S\n' % (i, i)
                           for i in range(1, 13)))
            imported = client.post('/api/gaia/model-mappings', data={
                'file': (io.BytesIO(csv.encode()), 'mapping.csv'),
            }, content_type='multipart/form-data')
            self.assertEqual(201, imported.status_code, imported.get_json())
            payload = {
                'shop_id': '5576543', 'source_path': str(template),
                'mapping_version_id': imported.get_json()['mapping']['id'],
                'selected_sku_ids': ['sku-%s' % i for i in range(1, 13)],
                'quality_labels': ['99新'], 'scene': 2, 'create_count': 1,
            }
            preview_response = client.post('/api/gaia/batches/preview', json=payload)
            self.assertEqual(200, preview_response.status_code, preview_response.get_json())
            preview = preview_response.get_json()['preview']
            self.assertEqual(1, len(calls))
            self.assertEqual(3, preview['file_count'])
            self.assertEqual([4, 4, 4], [item['row_count'] for item in preview['files']])
            self.assertEqual([1, 2, 3], [item['chunk_index'] for item in preview['files']])
            self.assertEqual([1007], [item['total_count'] for item in preview['live_goods_snapshot']])

            missing_token = client.post('/api/gaia/batches/generate', json={
                **payload, 'expected_source_sha256': preview['source_sha256'],
                'expected_quality_ids': preview['quality_ids'],
                'expected_split_plan_sha256': preview['split_plan_sha256'],
            })
            self.assertEqual(400, missing_token.status_code)
            self.assertIn('实时货品数量预览', missing_token.get_json()['message'])

            generated = client.post('/api/gaia/batches/generate', json={
                **payload, 'expected_source_sha256': preview['source_sha256'],
                'expected_quality_ids': preview['quality_ids'],
                'expected_inventory_probe_token': preview['inventory_probe_token'],
                'expected_split_plan_sha256': preview['split_plan_sha256'],
            })
            self.assertEqual(201, generated.status_code, generated.get_json())
            batch = generated.get_json()['batch']
            repository = GaiaBatchRepository(root_path / 'panel_data/gaia_listing/batches')
            self.assertEqual(3, len(batch['summary']['split_batch_ids']))
            self.assertEqual(3, len(batch['summary']['split_total_products']))
            self.assertEqual(['华为 Mate X6'] * 3,
                             [item['model'] for item in batch['summary']['split_total_products']])
            self.assertEqual([4, 4, 4], [
                repository.get(batch_id)['summary']['generated_rows']
                for batch_id in batch['summary']['split_batch_ids']])

    def test_existing_mapping_version_derives_goods_coverage_once(self):
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / 'mapping.csv'
            source.write_text(
                '品牌,型号,拍拍标品skuId,标品名称,货品ID,双维成色\n'
                '小米,小米 18 Pro,111,小米 18 Pro 国行,4830368,99S\n',
                encoding='utf-8',
            )
            repository = ModelMappingRepository(Path(root) / 'versions')
            metadata = repository.save_path(source)
            metadata_path = repository.root / (metadata['id'] + '.json')
            old_metadata = json.loads(metadata_path.read_text(encoding='utf-8'))
            old_metadata.pop('quality_coverage_version')
            metadata_path.write_text(json.dumps(old_metadata), encoding='utf-8')
            _, first = repository.load_records(metadata['id'])
            self.assertEqual(['4830368'], first[0].goods_for_quality('99'))
            sidecar = repository.root / (metadata['id'] + '.goods.json')
            self.assertTrue(sidecar.is_file())
            stale = json.loads(sidecar.read_text(encoding='utf-8'))
            stale['quality_coverage_version'] = 1
            stale['goods_by_sku']['111'] = {}
            sidecar.write_text(json.dumps(stale), encoding='utf-8')
            _, cached = ModelMappingRepository(repository.root).load_records(metadata['id'])
            self.assertEqual(first[0].goods_by_quality, cached[0].goods_by_quality)
            self.assertEqual(2, json.loads(sidecar.read_text(encoding='utf-8'))['quality_coverage_version'])

    def test_import_collects_all_goods_for_each_sku_and_grade(self):
        rows = [
            ['类目', '品牌', '型号', '货品ID', '拍拍标品skuId', '标品名称', '一维成色', '双维成色'],
            ['手机', '小米', '小米 18 Pro', '4830368', '111', '小米 18 Pro 国行', '', '99S'],
            ['手机', '小米', '小米 18 Pro', '4830369', '111', '小米 18 Pro 国行', '', '99B'],
            ['手机', '小米', '小米 18 Pro', '4830370', '111', '小米 18 Pro 国行', '', '95S'],
            ['手机', '小米', '小米 18 Pro', '4830371', '222', '小米 18 Pro 国行', '', '准新机'],
            ['手机', '小米', '小米 18 Pro', '4830372', '222', '小米 18 Pro 国行', '', '全新仅拆封'],
        ]
        report = _records_from_rows(rows)
        self.assertEqual(2, len(report.records))
        self.assertEqual({'99': ['4830368', '4830369'], '95': ['4830370']},
                         report.records[0].goods_by_quality)
        self.assertEqual({'100': ['4830371', '4830372']}, report.records[1].goods_by_quality)
        self.assertEqual([], report.records[0].goods_for_quality('100'))

    def test_preview_and_file_skip_only_unavailable_grade_combinations(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            config = root_path / 'config'
            config.mkdir()
            project = Path(__file__).resolve().parents[1]
            for name in ('gaia_excel_profiles.json', 'gaia_quality_mapping.json'):
                (config / name).write_bytes((project / 'config' / name).read_bytes())
            template = root_path / 'panel_data/gaia_listing/imports/template.xlsx'
            template.parent.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            workbook.save(template)
            app = Flask('gaia_goods_coverage')
            app.register_blueprint(create_gaia_blueprint(root, ['5576543']))
            client = app.test_client()
            csv = ('类目,品牌,型号,货品ID,拍拍标品skuId,标品名称,一维成色,双维成色\n'
                   '手机,小米,小米 18 Pro,4830368,111,小米 18 Pro 国行,,99S\n'
                   '手机,小米,小米 18 Pro,4830371,222,小米 18 Pro 国行,,全新仅拆封\n')
            imported = client.post('/api/gaia/model-mappings', data={
                'file': (io.BytesIO(csv.encode()), 'mapping.csv'),
            }, content_type='multipart/form-data')
            self.assertEqual(201, imported.status_code, imported.get_json())
            mapping_id = imported.get_json()['mapping']['id']
            payload = {'shop_id': '5576543', 'source_path': str(template),
                       'mapping_version_id': mapping_id, 'selected_sku_ids': ['111', '222'],
                       'quality_labels': ['99新', '准新机'], 'scene': 2, 'create_count': 1}
            preview_response = client.post('/api/gaia/batches/preview', json=payload)
            self.assertEqual(200, preview_response.status_code, preview_response.get_json())
            preview = preview_response.get_json()['preview']
            self.assertEqual([('111', '99'), ('222', '100')], [
                (row['values'][3], row['values'][4]) for row in preview['rows']])
            self.assertEqual([('111', '100'), ('222', '99')], [
                (item['standard_sku_id'], item['quality_id'])
                for item in preview['skipped_goods_combinations']])
            generated = client.post('/api/gaia/batches/generate', json={
                **payload, 'expected_source_sha256': preview['source_sha256'],
                'expected_quality_ids': preview['quality_ids'],
            })
            self.assertEqual(201, generated.status_code, generated.get_json())
            batch = generated.get_json()['batch']
            self.assertEqual(2, batch['summary']['skipped_goods_count'])
            book = load_workbook(batch['output_path'], read_only=True, data_only=True)
            try:
                self.assertEqual([('111', '99'), ('222', '100')], [
                    (str(row[3]), str(row[4])) for row in book.active.iter_rows(
                        min_row=2, min_col=1, max_col=5, values_only=True)])
            finally:
                book.close()
            missing_only = client.post('/api/gaia/batches/preview', json={
                **payload, 'selected_sku_ids': ['111'], 'quality_labels': ['准新机']})
            self.assertEqual(400, missing_only.status_code)
            self.assertIn('均无可确认', missing_only.get_json()['message'])

    def test_preview_assigns_one_aggregate_mark_per_model_after_skips(self):
        with tempfile.TemporaryDirectory() as root:
            project = Path(__file__).resolve().parents[1]
            mapping_path = Path(root) / 'mapping.csv'
            mapping_path.write_text(
                '类目,品牌,型号,货品ID,拍拍标品skuId,标品名称,双维成色\n'
                '手机,苹果,iPhone 17 Pro,11,sku-a1,苹果 iPhone 17 Pro 国行,99S\n'
                '手机,苹果,iPhone 17 Pro,12,sku-a2,苹果 iPhone 17 Pro 国行,99S\n'
                '手机,OPPO,Find X8,13,sku-b,OPPO Find X8 国行,95S\n'
                '手机,小米,Redmi K80,14,sku-c,小米 Redmi K80 国行,99S\n',
                encoding='utf-8',
            )
            mapping_repo = ModelMappingRepository(Path(root) / 'versions')
            mapping = mapping_repo.save_path(mapping_path)
            template = Path(root) / 'template.xlsx'
            workbook = Workbook()
            workbook.active.append(list(GAIA_TEMPLATE_HEADERS))
            workbook.save(template)
            prepared = prepare_gaia_template(
                source_path=template, mapping_version_id=mapping['id'],
                quality_label=['99新'], scene=2,
                selected_sku_ids=['sku-a1', 'sku-a2', 'sku-b', 'sku-c'],
                profiles=load_profiles(str(project / 'config/gaia_excel_profiles.json')),
                mapping_repository=mapping_repo,
                quality_repository=QualityMappingRepository(
                    str(project / 'config/gaia_quality_mapping.json')),
            )
            self.assertEqual([row.values()[1] for row in prepared.rows], ['A', 'A', 'B'])
            self.assertEqual([row['values'][1] for row in prepared.preview()['rows']], ['A', 'A', 'B'])


if __name__ == '__main__':
    unittest.main()
