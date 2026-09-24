import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.gaia_listing.application.template_generation import build_template_rows
from app.gaia_listing.domain.mapping_contract import GaiaMappingContractError, ModelMappingRecord
from app.gaia_listing.domain.template_contract import GAIA_TEMPLATE_HEADERS
from app.gaia_listing.infrastructure.model_mapping_import import import_model_mapping
from app.gaia_listing.infrastructure.quality_mapping_repository import QualityMappingRepository
from app.gaia_listing.infrastructure.template_writer import (
    GaiaTemplateWriteError,
    write_generated_template,
)


class GaiaTemplateGenerationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _template(self, filename="template.xlsx"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        for column, value in enumerate(GAIA_TEMPLATE_HEADERS, start=1):
            sheet.cell(1, column, value=value)
        path = self.root / filename
        workbook.save(path)
        return path

    def _quality(self, dy_quality_id="quality-99"):
        path = self.root / "quality.json"
        path.write_text(json.dumps({"qualities": [{"label": "99新", "dy_quality_id": dy_quality_id}]}), encoding="utf-8")
        return QualityMappingRepository(path).get_required("99新")

    def test_mapping_import_reads_operator_xlsx(self):
        path = self.root / "models.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["品牌名称", "型号名称", "标品ID（standardSkuId）", "品类"])
        sheet.append(["荣耀", "400 Pro", "sku-400", "手机"])
        workbook.save(path)
        records = import_model_mapping(path)
        self.assertEqual(records, [ModelMappingRecord("荣耀", "400 Pro", "sku-400", "手机")])

    def test_goods_export_deduplicates_same_sku_but_allows_model_variants(self):
        path = self.root / "goods.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["类目", "品牌", "型号", "拍拍标品skuId"])
        sheet.append(["手机", "荣耀", "400 Pro", "sku-a"])
        sheet.append(["手机", "荣耀", "400 Pro", "sku-a"])
        sheet.append(["手机", "荣耀", "400 Pro", "sku-b"])
        workbook.save(path)
        self.assertEqual(
            import_model_mapping(path),
            [
                ModelMappingRecord("荣耀", "400 Pro", "sku-a", "手机"),
                ModelMappingRecord("荣耀", "400 Pro", "sku-b", "手机"),
            ],
        )

    def test_goods_export_excludes_rows_without_explicit_mainland_evidence(self):
        path = self.root / "goods-region.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["类目", "品牌", "型号", "拍拍标品skuId", "标品名称"])
        sheet.append(["手机", "荣耀", "Magic V5", "sku-mainland", "荣耀 Magic V5 国行 16G+512G"])
        sheet.append(["手机", "荣耀", "Magic V5", "sku-other", "荣耀 Magic V5 其他地区 16G+512G"])
        sheet.append(["手机", "OPPO", "A3", "sku-unlabelled", "OPPO A3 8G+256G"])
        sheet.append(["手机", "OPPO", "A3", "sku-explicit-negative", "OPPO A3 非国行 8G+256G"])
        workbook.save(path)
        self.assertEqual(
            import_model_mapping(path),
            [ModelMappingRecord("荣耀", "Magic V5", "sku-mainland", "手机")],
        )

    def test_mapping_import_ignores_incorrect_a1_sheet_dimension(self):
        original = self.root / "original.xlsx"
        path = self.root / "bad-dimension.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["类目", "品牌", "型号", "拍拍标品skuId"])
        sheet.append(["手机", "荣耀", "400 Pro", "sku-a"])
        workbook.save(original)
        with zipfile.ZipFile(original, "r") as source, zipfile.ZipFile(path, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'<dimension ref="A1:D2"/>', b'<dimension ref="A1"/>')
                target.writestr(item, data)
        self.assertEqual(import_model_mapping(path)[0].standard_sku_id, "sku-a")

    def test_empty_quality_id_blocks_generation(self):
        path = self.root / "quality.json"
        path.write_text(json.dumps({"qualities": [{"label": "99新", "dy_quality_id": ""}]}), encoding="utf-8")
        with self.assertRaises(GaiaMappingContractError):
            QualityMappingRepository(path).get_required("99新")

    def test_writer_copies_template_and_writes_contract_columns(self):
        source = self._template()
        quality = self._quality()
        rows = build_template_rows(
            scene=2,
            models=[ModelMappingRecord("荣耀", "400 Pro", "sku-400", "手机")],
            quality=quality,
        )
        output = self.root / "output" / "gaia_upload.xlsx"
        write_generated_template(template_path=str(source), output_path=str(output), rows=rows)
        original = load_workbook(source)["Sheet1"]
        generated = load_workbook(output)["Sheet1"]
        self.assertEqual([original.cell(2, column).value for column in range(1, 9)], [None] * 8)
        self.assertEqual(
            [generated.cell(2, column).value for column in range(1, 9)],
            ["2", "A", "1", "sku-400", "quality-99", None, "D且E且G", "a"],
        )

    def test_writer_refuses_overwrite_of_template_data(self):
        source = self._template()
        workbook = load_workbook(source)
        workbook["Sheet1"].cell(2, 1, value="already-data")
        workbook.save(source)
        with self.assertRaises(GaiaTemplateWriteError):
            write_generated_template(
                template_path=str(source),
                output_path=str(self.root / "out.xlsx"),
                rows=build_template_rows(
                    scene=2,
                    models=[ModelMappingRecord("荣耀", "400 Pro", "sku-400", "手机")],
                    quality=self._quality(),
                ),
            )


if __name__ == "__main__":
    unittest.main()
