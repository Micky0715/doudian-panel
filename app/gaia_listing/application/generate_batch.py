"""Validated Gaia template preview and generation; never submits to Gaia."""

from __future__ import annotations

import os
import hashlib
import json
import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from ..domain.template_contract import GAIA_TEMPLATE_HEADERS, GaiaTemplateRow, aggregate_mark_for_index
from ..domain.states import GaiaBatchState
from ..infrastructure.excel_inspector import inspect_xlsx
from ..infrastructure.template_writer import GaiaTemplateWriteError, write_generated_template
from .create_batch import create_inspected_batch, sha256_file


@dataclass(frozen=True)
class PreparedGaiaTemplate:
    inspection: dict
    mapping: dict
    models: list
    rows: list
    quality_label: str
    quality_id: str
    source_hash: str
    header_row: int
    quality_labels: tuple[str, ...] = ()
    quality_ids: tuple[str, ...] = ()
    skipped_goods_combinations: tuple[dict, ...] = ()
    quality_link: str = "1"
    parts: tuple[tuple[int, int], ...] = ()
    part_metadata: tuple[dict, ...] = ()
    live_goods_snapshot: tuple[dict, ...] = ()
    split_plan_sha256: str = ""

    def preview(self, limit: int = 20) -> dict:
        return {
            "row_count": len(self.rows),
            "file_count": len(self.parts) or 1,
            "files": [
                {
                    "index": index,
                    "row_count": end - start,
                    "models": list(dict.fromkeys(
                        model.model for model in self.models[start:end])),
                    **(self.part_metadata[index - 1]
                       if index <= len(self.part_metadata) else {}),
                }
                for index, (start, end) in enumerate(self.parts, start=1)
            ],
            "mapping_version_id": self.mapping["id"],
            "mapping_sha256": self.mapping["sha256"],
            "quality_label": self.quality_label,
            "quality_id": self.quality_id,
            "quality_labels": list(self.quality_labels or (self.quality_label,)),
            "quality_ids": list(self.quality_ids or (self.quality_id,)),
            "quality_link": self.quality_link,
            "skipped_goods_combinations": list(self.skipped_goods_combinations),
            "skipped_goods_count": len(self.skipped_goods_combinations),
            "source_sha256": self.source_hash,
            "live_goods_snapshot": list(self.live_goods_snapshot),
            "split_plan_sha256": self.split_plan_sha256,
            "columns": list(GAIA_TEMPLATE_HEADERS),
            "rows": [
                {"excel_row": index + self.header_row + 1, "brand": model.brand, "model": model.model,
                 "category": model.category, "values": list(row.values())}
                for index, (model, row) in enumerate(zip(self.models[:limit], self.rows[:limit]))
            ],
            "truncated": len(self.rows) > limit,
        }


def split_model_parts(models: list, capacity: int) -> tuple[tuple[int, int], ...]:
    """Keep each model in one file; never mix models in a Gaia upload.

    Gaia has returned a 500-row rejection for verified 90-row workbooks that
    contain multiple models. The server's effective count is therefore not
    established by worksheet rows alone. One-model files are a conservative
    mitigation until the upstream counting rule is confirmed.
    """
    if capacity < 1:
        raise ValueError("盖亚模板表头超过单文件 500 行上限")
    groups: list[tuple[int, int]] = []
    start = 0
    while start < len(models):
        identity = (models[start].brand, models[start].model)
        end = start + 1
        while end < len(models) and (models[end].brand, models[end].model) == identity:
            end += 1
        if end - start > capacity:
            raise ValueError("%s %s 单型号有 %s 行，超过单文件 %s 行容量；为避免拆出重复商品，未自动上传" % (
                identity[0], identity[1], end - start, capacity))
        groups.append((start, end))
        start = end
    return tuple(groups)


LIVE_GOODS_SAFE_LIMIT = 400


def _model_key(brand: str, model: str) -> str:
    return "%s\u001f%s" % (str(brand).strip().casefold(), str(model).strip().casefold())


def _row_weight(model, row) -> int:
    """Use the imported goods mapping only to balance live-count chunks.

    The live Gaia total determines how many files are required.  The local
    mapping is not treated as the authoritative limit because it can lag Gaia.
    """
    goods = model.goods_for_quality(row.dy_quality_id)
    return max(1, len(goods)) if isinstance(goods, list) else 1


def _balanced_bins(indexes: list[int], weights: list[int], count: int) -> list[list[int]]:
    if count < 1 or count > len(indexes):
        raise ValueError("盖亚实时货品量需要拆成 %s 份，但当前只有 %s 个可拆标品/成色行" % (
            count, len(indexes)))
    bins = [[] for _ in range(count)]
    totals = [0] * count
    # Largest-first balancing avoids concentrating high-coverage SKUs in one
    # file.  Restore source order within each file for reproducibility.
    for index, weight in sorted(zip(indexes, weights), key=lambda item: (-item[1], item[0])):
        target = min(range(count), key=lambda slot: (totals[slot], len(bins[slot]), slot))
        bins[target].append(index)
        totals[target] += weight
    return [sorted(group) for group in bins]


def apply_live_goods_plan(prepared: PreparedGaiaTemplate, snapshot: dict,
                          safe_limit: int = LIVE_GOODS_SAFE_LIMIT) -> PreparedGaiaTemplate:
    """Split by Gaia's current model goods totals before any upload occurs.

    A workbook row can expand to many Gaia goods.  Therefore the worksheet row
    count is only a hard file-format guard; it is not the upstream 500-item
    count.  Every model is kept isolated and may produce multiple products.
    """
    if safe_limit < 1 or safe_limit >= 500:
        raise ValueError("盖亚实时货品安全阈值必须在 1 到 499 之间")
    results = snapshot.get("models") if isinstance(snapshot, dict) else None
    if not isinstance(results, list):
        raise ValueError("盖亚实时货品数量返回格式无效，未生成上传文件")
    by_key = {}
    for item in results:
        if not isinstance(item, dict):
            raise ValueError("盖亚实时货品数量包含无效记录，未生成上传文件")
        brand, model = str(item.get("brand") or "").strip(), str(item.get("model") or "").strip()
        count = item.get("total_count")
        if not brand or not model or not isinstance(count, int) or isinstance(count, bool) or count < 0:
            raise ValueError("盖亚实时货品数量缺少完整型号或有效数量，未生成上传文件")
        key = _model_key(brand, model)
        if key in by_key:
            raise ValueError("盖亚实时货品数量重复返回 %s %s，未生成上传文件" % (brand, model))
        by_key[key] = dict(item, brand=brand, model=model, total_count=count)

    grouped: dict[str, list[int]] = {}
    identities: dict[str, tuple[str, str]] = {}
    for index, model in enumerate(prepared.models):
        key = _model_key(model.brand, model.model)
        grouped.setdefault(key, []).append(index)
        identities[key] = (model.brand, model.model)
    missing = [identities[key] for key in grouped if key not in by_key]
    extra = [item for key, item in by_key.items() if key not in grouped]
    if missing or extra:
        detail = ("缺少：" + "、".join("%s %s" % item for item in missing)) if missing else (
            "多出：" + "、".join("%s %s" % (item["brand"], item["model"]) for item in extra))
        raise ValueError("盖亚实时货品数量与本次完整型号不一致（%s），未生成上传文件" % detail)

    ordered_rows, ordered_models, parts, metadata = [], [], [], []
    for key, indexes in grouped.items():
        brand, model_name = identities[key]
        total_count = by_key[key]["total_count"]
        chunk_total = max(1, math.ceil(total_count / safe_limit))
        weights = [_row_weight(prepared.models[index], prepared.rows[index]) for index in indexes]
        bins = _balanced_bins(indexes, weights, chunk_total)
        for chunk_index, selected in enumerate(bins, start=1):
            start = len(ordered_rows)
            ordered_rows.extend(prepared.rows[index] for index in selected)
            ordered_models.extend(prepared.models[index] for index in selected)
            end = len(ordered_rows)
            if end - start > 500 - prepared.header_row:
                raise ValueError("%s %s 第 %s 份仍有 %s 行，超过 Excel 单文件容量，未生成上传文件" % (
                    brand, model_name, chunk_index, end - start))
            parts.append((start, end))
            metadata.append({
                "brand": brand,
                "model": model_name,
                "chunk_index": chunk_index,
                "chunk_total": chunk_total,
                "live_model_goods_total": total_count,
                "live_goods_safe_limit": safe_limit,
                "local_weight": sum(_row_weight(prepared.models[index], prepared.rows[index])
                                    for index in selected),
            })

    normalized_snapshot = tuple(dict(item) for item in results)
    fingerprint_payload = {
        "source_sha256": prepared.source_hash,
        "mapping_sha256": prepared.mapping["sha256"],
        "quality_ids": list(prepared.quality_ids),
        "safe_limit": safe_limit,
        "models": list(normalized_snapshot),
        "parts": metadata,
    }
    fingerprint = hashlib.sha256(json.dumps(
        fingerprint_payload, ensure_ascii=False, sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    return PreparedGaiaTemplate(
        inspection=prepared.inspection, mapping=prepared.mapping,
        models=ordered_models, rows=ordered_rows,
        quality_label=prepared.quality_label, quality_id=prepared.quality_id,
        source_hash=prepared.source_hash, header_row=prepared.header_row,
        quality_labels=prepared.quality_labels, quality_ids=prepared.quality_ids,
        skipped_goods_combinations=prepared.skipped_goods_combinations,
        quality_link=prepared.quality_link, parts=tuple(parts),
        part_metadata=tuple(metadata), live_goods_snapshot=normalized_snapshot,
        split_plan_sha256=fingerprint,
    )


def prepare_gaia_template(*, source_path, mapping_version_id, quality_label, scene,
                          selected_sku_ids,
                          profiles, mapping_repository, quality_repository,
                          rejected_variants=None, quality_link="1") -> PreparedGaiaTemplate:
    if str(scene) not in {"1", "2"}:
        raise ValueError("场域只能为 1（直播间）或 2（直播货架）")
    if str(quality_link) not in {"0", "1"}:
        raise ValueError("是否质检链接只能选择是（1）或否（0）")
    quality_link = str(quality_link)
    inspection = inspect_xlsx(source_path, profiles)
    matched = inspection["profile_matches"]
    if inspection["status"] != "recognized" or {item["id"] for item in matched} != {"gaia_publish_template"}:
        raise ValueError("文件不是唯一识别的盖亚双维成色发品模板")
    mapping, catalogue = mapping_repository.load_records(mapping_version_id)
    labels = quality_label if isinstance(quality_label, list) else [quality_label]
    if not labels or any(not isinstance(label, str) or not label.strip() for label in labels):
        raise ValueError("请至少选择一种成色")
    labels = [label.strip() for label in labels]
    if len(labels) != len(set(labels)):
        raise ValueError("成色不能重复选择")
    qualities = [quality_repository.get_required(label) for label in labels]
    if len(catalogue) != mapping.get("record_count"):
        raise ValueError("型号映射版本行数不一致，请重新上传")
    if not isinstance(selected_sku_ids, list):
        raise ValueError("请提供本次待生成的标品 SKU ID 清单")
    requested = [str(value or "").strip() for value in selected_sku_ids]
    if not requested or any(not value for value in requested):
        raise ValueError("本次待生成的标品 SKU ID 不能为空")
    if len(requested) != len(set(requested)):
        raise ValueError("本次标品 SKU ID 清单存在重复")
    if len(requested) * len(qualities) > 2000:
        raise ValueError("单次任务最多检查 2000 行标品 SKU 与成色组合")
    by_sku = {item.standard_sku_id: item for item in catalogue}
    missing = [value for value in requested if value not in by_sku]
    if missing:
        sample = "、".join(missing[:10])
        raise ValueError("以下标品 SKU ID 不在当前映射版本中：%s%s" % (
            sample, " 等" if len(missing) > 10 else "",
        ))
    selected_models = [by_sku[value] for value in requested]
    # Keep all SKUs of one complete model together even if an API caller sent
    # interleaved IDs. Splitting the same model across files may create two
    # products and makes post-submit reconciliation ambiguous.
    model_groups = {}
    for model in selected_models:
        model_groups.setdefault((model.brand, model.model), []).append(model)
    selected_models = [model for group in model_groups.values() for model in group]
    if mapping.get("mainland_filter_applied") and any(
            model.goods_by_quality is None for model in selected_models):
        raise ValueError("映射表缺少货品ID或双维成色列，不能确认标品与所选成色的货品关系；请更新官方原表")
    rejected_variants = rejected_variants or {}
    rows, models, skipped = [], [], []
    aggregate_marks = {}
    for model in selected_models:
        for quality in qualities:
            sku_id, quality_id = model.standard_sku_id, quality.dy_quality_id
            available = model.goods_for_quality(quality_id)
            rejected_batch = rejected_variants.get((sku_id, quality_id))
            if rejected_batch or available == []:
                skipped.append({
                    "standard_sku_id": sku_id, "quality_label": quality.label,
                    "quality_id": quality_id,
                    "reason": "gaia_rejected" if rejected_batch else "no_matching_goods",
                    "rejected_batch_id": rejected_batch,
                })
                continue
            model_key = (model.brand, model.model)
            if model_key not in aggregate_marks:
                aggregate_marks[model_key] = aggregate_mark_for_index(len(aggregate_marks))
            rows.append(GaiaTemplateRow(
                scene=str(scene), standard_sku_id=sku_id,
                dy_quality_id=quality_id, category=model.category,
                aggregate_mark=aggregate_marks[model_key],
                quality_link=quality_link,
            ))
            models.append(model)
    if not rows:
        examples = "、".join(
            "%s/%s" % (item["standard_sku_id"], item["quality_label"])
            for item in skipped[:10]
        )
        raise ValueError(
            "所选标品与成色均无可确认的对应货品，未生成或上传；"
            "请核对原表双维成色及盖亚拒绝记录。跳过组合：%s%s" % (
                examples, " 等" if len(skipped) > 10 else "",
            )
        )
    source_hash = sha256_file(source_path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        header_row = next((index for index, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=10, max_col=8, values_only=True), start=1)
            if tuple(values) == GAIA_TEMPLATE_HEADERS), 0)
        if header_row and any(any(value not in (None, "") for value in values)
                              for values in sheet.iter_rows(min_row=header_row + 1, max_col=8, values_only=True)):
            raise GaiaTemplateWriteError("模板已有 A:H 数据，拒绝覆盖；请重新下载空模板")
    finally:
        workbook.close()
    if not header_row:
        raise GaiaTemplateWriteError("模板未找到完整 A:H 表头")
    parts = split_model_parts(models, 500 - header_row)
    quality_labels = tuple(item.label for item in qualities)
    quality_ids = tuple(item.dy_quality_id for item in qualities)
    return PreparedGaiaTemplate(
        inspection, mapping, models, rows,
        quality_labels[0] if len(quality_labels) == 1 else "、".join(quality_labels),
        quality_ids[0] if len(quality_ids) == 1 else ",".join(quality_ids),
        source_hash, header_row, quality_labels, quality_ids, tuple(skipped), quality_link, parts,
    )


def _verify_output(path: Path, rows: list) -> None:
    """Read back generated A:H data before marking the batch validated."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(min_row=1, max_col=8, values_only=True)
        headers = ()
        for _ in range(10):
            headers = tuple(next(iterator, ()))
            if headers == GAIA_TEMPLATE_HEADERS:
                break
        if headers != GAIA_TEMPLATE_HEADERS:
            raise GaiaTemplateWriteError("输出文件表头不一致")
        for index, expected in enumerate(rows, start=2):
            actual = tuple(str(value or "") for value in next(iterator, ()))
            if actual != expected.values():
                raise GaiaTemplateWriteError("输出文件第 %s 行回读不一致" % index)
        if any(value not in (None, "") for value in next(iterator, ())):
            raise GaiaTemplateWriteError("输出文件存在未预期的额外数据行")
    finally:
        workbook.close()


def generate_gaia_batch(*, repository, outputs_dir, shop_id, scene, create_count,
                        source_path, profiles, prepared: PreparedGaiaTemplate,
                        selection_report=None):
    if int(create_count) != 1:
        raise ValueError("本阶段每行创建数量固定为 1")
    parts = prepared.parts or ((0, len(prepared.rows)),)
    batches = []
    for part_index, (start, end) in enumerate(parts, start=1):
        batches.append(_generate_one_batch(
            repository=repository, outputs_dir=outputs_dir, shop_id=shop_id,
            scene=scene, create_count=create_count, source_path=source_path,
            profiles=profiles, prepared=prepared, selection_report=selection_report,
            part_index=part_index, part_total=len(parts), start=start, end=end,
        ))
    if len(batches) > 1:
        ids = [batch.id for batch in batches]
        # Equal brand/model pairs are intentionally preserved when a single
        # high-volume model was split into several Gaia products.
        total_products = []
        for part_index, (start, _end) in enumerate(parts, start=1):
            model = prepared.models[start]
            details = (prepared.part_metadata[part_index - 1]
                       if part_index <= len(prepared.part_metadata) else {})
            total_products.append({
                "brand": model.brand,
                "model": model.model,
                "chunk_index": details.get("chunk_index", 1),
                "chunk_total": details.get("chunk_total", 1),
            })
        for batch in batches:
            batch.summary.update({
                "split_batch_ids": ids,
                "split_root_batch_id": ids[0],
                "split_total_rows": len(prepared.rows),
                "split_total_products": total_products,
                "split_plan_sha256": prepared.split_plan_sha256,
                "live_goods_snapshot": list(prepared.live_goods_snapshot),
            })
            repository.save(batch)
    return batches[0]


def _generate_one_batch(*, repository, outputs_dir, shop_id, scene, create_count,
                        source_path, profiles, prepared: PreparedGaiaTemplate,
                        selection_report, part_index: int, part_total: int,
                        start: int, end: int):
    rows, models = prepared.rows[start:end], prepared.models[start:end]
    batch, _ = create_inspected_batch(
        repository, shop_id=shop_id, scene=scene, create_count=create_count,
        source_path=source_path, profiles=profiles,
    )
    batch.mapping_version_id = prepared.mapping["id"]
    batch.mapping_hash = prepared.mapping["sha256"]
    batch.quality_label = prepared.quality_label
    batch.quality_id = prepared.quality_id
    batch.quality_labels = list(prepared.quality_labels or (prepared.quality_label,))
    batch.quality_ids = list(prepared.quality_ids or (prepared.quality_id,))
    output = Path(outputs_dir) / batch.id / "gaia_upload.xlsx"
    temporary = output.with_suffix(".tmp.xlsx")
    try:
        if batch.source_hash != prepared.source_hash or sha256_file(source_path) != batch.source_hash:
            raise GaiaTemplateWriteError("源模板在校验后发生变化")
        write_generated_template(template_path=source_path, output_path=str(temporary), rows=rows)
        _verify_output(temporary, rows)
        if sha256_file(source_path) != batch.source_hash:
            raise GaiaTemplateWriteError("源模板在生成期间发生变化")
        os.replace(temporary, output)
        batch.output_path = os.path.realpath(output)
        batch.output_hash = sha256_file(output)
        batch.state = GaiaBatchState.VALIDATION_PASSED.value
        products = list(dict.fromkeys((model.brand, model.model) for model in models))
        part_details = (prepared.part_metadata[part_index - 1]
                        if part_index <= len(prepared.part_metadata) else {})
        batch.summary = {
            "generated_rows": len(rows),
            "split_part_index": part_index,
            "split_part_total": part_total,
            "quality_link": prepared.quality_link,
            "skipped_goods_count": len(prepared.skipped_goods_combinations),
            "skipped_goods_combinations": list(prepared.skipped_goods_combinations),
            "expected_product_count": len(products),
            "expected_products": [{"brand": brand, "model": model} for brand, model in products],
            "split_plan_sha256": prepared.split_plan_sha256,
            "live_goods_snapshot": list(prepared.live_goods_snapshot),
            **part_details,
            "blocking_errors": 0,
        }
        if selection_report is not None:
            batch.summary.update({
                "input_model_count": len(selection_report["results"]),
                "matched_model_count": selection_report["matched_model_count"],
                "skipped_models": selection_report["skipped_models"],
                "excluded_skus": selection_report["excluded_skus"],
                "partial_selection": selection_report["partial"],
            })
    except Exception as exc:
        batch.state = GaiaBatchState.VALIDATION_FAILED.value
        batch.errors = [{"code": "template_generation_failed", "message": str(exc)}]
        if temporary.exists():
            temporary.unlink()
        repository.save(batch)
        raise
    repository.save(batch)
    return batch
