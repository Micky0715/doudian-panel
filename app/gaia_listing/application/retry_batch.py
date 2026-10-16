"""Create a fresh, auditable batch after an atomic Gaia goodsId rejection."""

import os
from pathlib import Path

from openpyxl import load_workbook

from ..domain.states import GaiaBatchState
from ..domain.template_contract import GAIA_TEMPLATE_HEADERS, GaiaTemplateRow
from ..infrastructure.template_writer import write_generated_template
from .create_batch import create_inspected_batch, sha256_file
from .generate_batch import _verify_output


def create_remaining_batch(*, previous, rejected_pair, repository, outputs_dir,
                           mapping_repository, profiles):
    """Exclude exactly one proven-bad pair; never edit or resubmit the old file."""
    if previous.get("state") != GaiaBatchState.GAIA_REJECTED.value:
        raise ValueError("只有已明确拒绝的批次可生成剩余组合")
    mapping, catalogue = mapping_repository.load_records(previous["mapping_version_id"])
    if mapping["sha256"] != previous["mapping_hash"]:
        raise ValueError("映射版本已变化")
    by_sku = {item.standard_sku_id: item for item in catalogue}
    workbook = load_workbook(previous["output_path"], read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(min_row=1, max_col=8, values_only=True)
        if tuple(next(values, ())) != GAIA_TEMPLATE_HEADERS:
            raise ValueError("原批次表头已变化")
        retained = []
        removed = 0
        seen = set()
        for raw in values:
            if not any(value not in (None, "") for value in raw):
                continue
            sku, quality = str(raw[3] or "").strip(), str(raw[4] or "").strip()
            model = by_sku.get(sku)
            if model is None or quality not in previous["quality_ids"]:
                raise ValueError("原批次存在无法核对的标品或成色")
            row = GaiaTemplateRow(
                str(previous["scene"]), sku, quality, model.category,
                aggregate_mark=str(raw[1] or ""), quality_link=str(raw[2] or ""),
            )
            if tuple(str(value or "") for value in raw) != row.values() or (sku, quality) in seen:
                raise ValueError("原批次行与受控模板不一致")
            seen.add((sku, quality))
            if (sku, quality) == tuple(rejected_pair):
                removed += 1
            else:
                retained.append(row)
    finally:
        workbook.close()
    if removed != 1:
        raise ValueError("盖亚拒绝的组合未在原批次中唯一出现")
    if not retained:
        return None

    batch, _ = create_inspected_batch(
        repository, shop_id=previous["shop_id"], scene=previous["scene"],
        create_count=1, source_path=previous["source_path"], profiles=profiles,
    )
    batch.mapping_version_id = mapping["id"]
    batch.mapping_hash = mapping["sha256"]
    for field in ("quality_label", "quality_id", "quality_labels", "quality_ids"):
        setattr(batch, field, previous[field])
    output = Path(outputs_dir) / batch.id / "gaia_upload.xlsx"
    temporary = output.with_suffix(".tmp.xlsx")
    if batch.source_hash != previous["source_hash"]:
        raise ValueError("源模板已变化")
    try:
        write_generated_template(template_path=batch.source_path,
                                 output_path=str(temporary), rows=retained)
        _verify_output(temporary, retained)
        if sha256_file(batch.source_path) != batch.source_hash:
            raise ValueError("源模板在生成期间变化")
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()
    batch.output_path = os.path.realpath(output)
    batch.output_hash = sha256_file(output)
    batch.state = GaiaBatchState.VALIDATION_PASSED.value
    products = list(dict.fromkeys((by_sku[row.standard_sku_id].brand,
                                   by_sku[row.standard_sku_id].model) for row in retained))
    batch.summary = {
        "generated_rows": len(retained), "expected_product_count": len(products),
        "expected_products": [{"brand": brand, "model": model} for brand, model in products],
        "blocking_errors": 0, "retry_of": previous["id"],
        "excluded_pair": list(rejected_pair),
    }
    repository.save(batch)
    return batch.to_dict()
