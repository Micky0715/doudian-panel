"""Read-only diagnosis of a Gaia per-row import result workbook.

An accepted upload is not a successful product creation.  This module never
changes a batch or retries a submission; the result workbook must be supplied
explicitly by an operator and is not treated as a platform status readback.
"""

from __future__ import annotations

from collections import Counter
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook


CASCADE_ERROR = "同组其他数据有错误"
RULE_PRICE_ERROR = "获取抖快渠道规则价格失败"


def _value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows(path: str | Path, columns: int):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        # Gaia result exports can incorrectly declare dimension A1:A1 while
        # containing a complete table.  Trust the worksheet stream instead.
        sheet.reset_dimensions()
        for number, row in enumerate(sheet.iter_rows(max_col=columns, values_only=True), start=1):
            if number > 5001:
                raise ValueError("结果表超过 5000 行，不在单批次诊断范围")
            yield tuple(_value(value) for value in row)
    finally:
        workbook.close()


def diagnose_gaia_result(*, batch: dict, result_path: str | Path,
                         goods_by_sku: dict | None = None) -> dict:
    """Correlate an explicitly supplied result with one immutable upload file.

    SKU/quality coverage is checked before attributing any result to a batch.
    Historical mapping coverage can explain a missing goods ID, but cannot
    prove the live Gaia rule-price configuration or product creation status.
    """
    output_path = batch.get("output_path")
    if not output_path or not Path(output_path).is_file():
        raise ValueError("批次上传文件不存在，不能核对结果表归属")
    expected_hash = batch.get("output_hash")
    if expected_hash and sha256(Path(output_path).read_bytes()).hexdigest() != expected_hash:
        raise ValueError("批次上传文件与生成时的哈希不一致，不能核对结果表")
    source_rows = list(_rows(output_path, 8))
    if len(source_rows) < 2:
        raise ValueError("批次上传文件没有数据行")
    if ("standardSkuId" not in source_rows[0][3]
            or "dyQualityId" not in source_rows[0][4]):
        raise ValueError("批次上传文件表头不符合盖亚模板")
    expected_pairs = {(row[3], row[4]) for row in source_rows[1:] if row[3] and row[4]}
    if not expected_pairs:
        raise ValueError("批次上传文件缺少标品或成色")

    result_rows = list(_rows(result_path, 9))
    if len(result_rows) < 2:
        raise ValueError("盖亚结果表没有数据行")
    if ("goodsId" not in result_rows[0][2] or result_rows[0][6] != "错误信息"
            or "standardSkuId" not in result_rows[0][7]
            or "dyQualityId" not in result_rows[0][8]):
        raise ValueError("结果表表头不符合盖亚创建优品结果格式")
    records = []
    for excel_row, row in enumerate(result_rows[1:], start=2):
        if not any(row):
            continue
        records.append({
            "excel_row": excel_row,
            "goods_id": row[2],
            "error": row[6],
            "standard_sku_id": row[7],
            "quality_id": row[8],
        })
    if not records:
        raise ValueError("盖亚结果表没有有效记录")
    observed_pairs = {(item["standard_sku_id"], item["quality_id"]) for item in records}
    if not observed_pairs <= expected_pairs or not all(all((sku, quality)) for sku, quality in observed_pairs):
        raise ValueError("结果表含本批次未上传的标品或成色，不能归属到该批次")
    expected_skus = {sku for sku, _ in expected_pairs}
    observed_skus = {sku for sku, _ in observed_pairs}
    if observed_skus != expected_skus:
        raise ValueError("结果表未覆盖本批次全部标品，不能唯一归属到该批次")

    primary = []
    cascades = 0
    for item in records:
        if item["error"] == CASCADE_ERROR:
            cascades += 1
        elif item["error"]:
            finding = dict(item)
            if RULE_PRICE_ERROR in item["error"]:
                finding["category"] = "rule_price_lookup_failed"
            else:
                finding["category"] = "other_platform_error"
            if goods_by_sku is not None and item["goods_id"]:
                known = (goods_by_sku.get(item["standard_sku_id"]) or {}).get(
                    item["quality_id"])
                finding["goods_in_mapping_snapshot"] = (
                    None if known is None else item["goods_id"] in {_value(value) for value in known}
                )
            primary.append(finding)

    rule_price_errors = [item for item in primary
                         if item["category"] == "rule_price_lookup_failed"]
    missing_pairs = expected_pairs - observed_pairs
    if rule_price_errors:
        next_step = ("先按首发错误的货品 ID 在盖亚核对抖快渠道规则价格；"
                     "若货品不在历史映射中，更新官方映射后再核对。"
                     "同组错误不是独立首发原因，已受理批次不得直接重导。")
    else:
        next_step = "按首发错误逐项核对盖亚结果；已受理批次不得直接重导。"

    return {
        "batch_id": batch.get("id"),
        "shop_id": batch.get("shop_id"),
        "result_rows": len(records),
        "result_unique_skus": len(observed_skus),
        "expected_sku_quality_pairs": len(expected_pairs),
        "observed_sku_quality_pairs": len(observed_pairs),
        "missing_sku_quality_pairs": len(missing_pairs),
        "pair_coverage": "complete" if not missing_pairs else "partial",
        "result_attribution": "sku_scope_only",
        "error_counts": dict(Counter(item["error"] for item in records if item["error"])),
        "primary_errors": primary,
        "cascade_error_count": cascades,
        "rule_price_error_count": len(rule_price_errors),
        "next_step": next_step,
        "platform_product_state": "unknown",
        "note": ("结果表未包含本地批次唯一标识，只能核对标品范围，不能证明结果表唯一归属该批次；"
                 "映射为历史快照，不能证明盖亚当前价格规则或抖店草稿状态。"),
    }
