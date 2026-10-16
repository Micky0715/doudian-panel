"""Import operator-managed model mapping workbooks without business side effects."""

from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass
from dataclasses import replace
from pathlib import Path
from typing import Iterable, List, Mapping

from openpyxl import load_workbook
from model_catalogue import canonical_brand, model_key

from ..domain.mapping_contract import GaiaMappingContractError, ModelMappingRecord


MAX_MAPPING_ROWS = 250_000

NON_MAINLAND_MARKERS = (
    "非国行", "其他地区", "港版", "香港", "澳门", "台版", "台湾",
    "美版", "日版", "韩版", "欧版", "澳版", "海外版", "国际版", "亚太版",
)


HEADER_ALIASES = {
    "brand": ("品牌", "品牌名称", "brand"),
    "model": ("型号", "型号名称", "机型", "model"),
    "standard_sku_id": ("标品id", "标品sku id", "standardskuid", "拍拍标品skuid"),
    "category": ("品类", "品类名称", "类目", "商品品类", "category"),
    "product_name": ("标品名称", "商品名称", "productname"),
    "goods_id": ("货品id", "goodsid"),
    "dual_quality": ("双维成色",),
}


def _dy_quality_id_from_export(value: object) -> str | None:
    """Only map the explicit grade families found in Gaia's goods export."""
    code = str(value or "").strip().upper()
    if code in {"准新机", "全新仅拆封"}:
        return "100"
    match = re.fullmatch(r"(99|95|90)[A-Z]", code)
    return match.group(1) if match else None


@dataclass(frozen=True)
class ModelMappingImportReport:
    """Auditable result of one read-only model-mapping import."""

    records: List[ModelMappingRecord]
    source_data_rows: int
    accepted_rows: int
    duplicate_rows: int
    mainland_filter_applied: bool
    excluded_non_mainland_rows: int
    excluded_non_mainland_sku_count: int


def _normal_header(value: object) -> str:
    return re.sub(r"[\W_]+", "", str(value or "").casefold())


def is_explicit_mainland_product_name(value: object) -> bool:
    """Accept only positive mainland evidence and reject contradictory labels."""
    name = re.sub(r"\s+", "", str(value or "")).casefold()
    return "国行" in name and not any(marker.casefold() in name for marker in NON_MAINLAND_MARKERS)


def _index_headers(headers: Iterable[object]) -> Mapping[str, int]:
    normal = [_normal_header(value) for value in headers]
    result = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            candidate = _normal_header(alias)
            try:
                result[field] = next(
                    index for index, header in enumerate(normal)
                    if header == candidate or candidate in header
                )
                break
            except StopIteration:
                continue
    missing = [field for field in ("brand", "model", "standard_sku_id") if field not in result]
    if missing:
        if any("聚合标识" in header for header in normal):
            raise GaiaMappingContractError(
                "选中的是双维成色发品模板，不是型号映射表；请上传直播货架货品价格规则数据原表"
            )
        raise GaiaMappingContractError("型号映射表缺少列：%s" % "、".join(missing))
    return result


def _records_from_rows(rows: Iterable[Iterable[object]]) -> ModelMappingImportReport:
    iterator = iter(rows)
    headers = next(iterator, None)
    if headers is None:
        raise GaiaMappingContractError("型号映射表没有表头")
    indexes = _index_headers(headers)
    records = []
    sku_owners = {}
    goods_by_sku: dict[str, dict[str, set[str]]] = {}
    has_goods_coverage = "goods_id" in indexes and "dual_quality" in indexes
    source_data_rows = 0
    accepted_rows = 0
    duplicate_rows = 0
    excluded_non_mainland_rows = 0
    excluded_non_mainland_skus = set()
    mainland_filter_applied = "product_name" in indexes
    for number, raw in enumerate(iterator, start=2):
        if number > MAX_MAPPING_ROWS + 1:
            raise GaiaMappingContractError("型号映射表超过 25 万行限制")
        values = list(raw)
        if not any(str(value or "").strip() for value in values):
            continue
        source_data_rows += 1
        def value(field: str) -> str:
            index = indexes.get(field)
            raw_value = values[index] if index is not None and index < len(values) else None
            if raw_value is None:
                return ""
            if isinstance(raw_value, float) and raw_value.is_integer():
                return str(int(raw_value))
            return str(raw_value).strip()
        # The Gaia goods-price export has no standalone market/region column.
        # Its product name is the only authoritative evidence available in the
        # workbook.  Fail closed: an unlabelled row is not treated as mainland.
        if mainland_filter_applied and not is_explicit_mainland_product_name(value("product_name")):
            excluded_non_mainland_rows += 1
            excluded_sku = value("standard_sku_id")
            if excluded_sku:
                excluded_non_mainland_skus.add(excluded_sku)
            continue
        try:
            record = ModelMappingRecord(
                brand=value("brand"),
                model=value("model"),
                standard_sku_id=value("standard_sku_id"),
                category=value("category") or "手机",
            )
        except GaiaMappingContractError as exc:
            raise GaiaMappingContractError("型号映射表第 %s 行：%s" % (number, exc)) from exc
        identity = (
            record.category.casefold(), canonical_brand(record.brand), model_key(record.model),
        )
        owner = sku_owners.get(record.standard_sku_id)
        if owner is not None:
            if owner != identity:
                raise GaiaMappingContractError(
                    "型号映射表同一标品ID对应多个型号或品类/品牌：%s" % record.standard_sku_id
                )
        if has_goods_coverage:
            quality_id = _dy_quality_id_from_export(value("dual_quality"))
            goods_id = value("goods_id")
            if quality_id and goods_id:
                goods_by_sku.setdefault(record.standard_sku_id, {}).setdefault(quality_id, set()).add(goods_id)
        if owner is not None:
            # 货品维度导出会因成色、价格规则等重复出现同一标品 SKU。
            # 映射版本只保留其第一次出现的位置，输出仍然稳定可审计。
            duplicate_rows += 1
            continue
        sku_owners[record.standard_sku_id] = identity
        records.append(record)
        accepted_rows += 1
    if has_goods_coverage:
        records = [replace(record, goods_by_quality={
            quality_id: sorted(goods_ids) for quality_id, goods_ids in
            sorted(goods_by_sku.get(record.standard_sku_id, {}).items())
        }) for record in records]
    if not records:
        if mainland_filter_applied and excluded_non_mainland_rows:
            raise GaiaMappingContractError("型号映射表没有明确标注“国行”的有效数据行")
        raise GaiaMappingContractError("型号映射表没有有效数据行")
    return ModelMappingImportReport(
        records=records,
        source_data_rows=source_data_rows,
        accepted_rows=accepted_rows,
        duplicate_rows=duplicate_rows,
        mainland_filter_applied=mainland_filter_applied,
        excluded_non_mainland_rows=excluded_non_mainland_rows,
        excluded_non_mainland_sku_count=len(excluded_non_mainland_skus),
    )


def import_model_mapping_with_report(path: str) -> ModelMappingImportReport:
    """Read one xlsx/csv mapping file; no output is written in this module."""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                # 部分盖亚大表把 worksheet dimension 错写成 A1；read-only
                # 模式会因此只迭代一个单元格。强制重新计算实际使用范围。
                sheet.reset_dimensions()
                return _records_from_rows(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        except (OSError, zipfile.BadZipFile) as exc:
            raise GaiaMappingContractError("无法读取型号映射表") from exc
    if suffix == ".csv":
        try:
            with source.open("r", encoding="utf-8-sig", newline="") as handle:
                return _records_from_rows(csv.reader(handle))
        except (OSError, UnicodeError) as exc:
            raise GaiaMappingContractError("无法读取型号映射表") from exc
    raise GaiaMappingContractError("型号映射表仅支持 .xlsx 或 .csv；.xls 请先另存为 .xlsx")


def import_model_mapping(path: str) -> List[ModelMappingRecord]:
    """Backward-compatible record-only interface used by existing callers."""
    return import_model_mapping_with_report(path).records
