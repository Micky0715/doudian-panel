"""Copy a recognized Gaia template and write only its A:H data range."""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook

from ..domain.template_contract import GAIA_TEMPLATE_HEADERS, GaiaTemplateRow


class GaiaTemplateWriteError(ValueError):
    pass


def _header_row(sheet) -> int:
    wanted = tuple(GAIA_TEMPLATE_HEADERS)
    for row_number in range(1, min(sheet.max_row, 10) + 1):
        values = tuple(str(sheet.cell(row_number, column).value or "").strip() for column in range(1, len(wanted) + 1))
        if values == wanted:
            return row_number
    raise GaiaTemplateWriteError("模板未找到完整 A:H 表头，拒绝写入")


def _assert_empty_data_area(sheet, first_data_row: int) -> None:
    for row_number in range(first_data_row, sheet.max_row + 1):
        if any(str(sheet.cell(row_number, column).value or "").strip() for column in range(1, 9)):
            raise GaiaTemplateWriteError("模板已有 A:H 数据，拒绝覆盖；请重新下载空模板")


def write_generated_template(*, template_path: str, output_path: str, rows: Iterable[GaiaTemplateRow]) -> str:
    """Write a new workbook; the source template path is never opened writable."""
    source = Path(template_path)
    destination = Path(output_path)
    if source.resolve() == destination.resolve():
        raise GaiaTemplateWriteError("输出文件不能覆盖原模板")
    materialized: Sequence[GaiaTemplateRow] = tuple(rows)
    if not materialized:
        raise GaiaTemplateWriteError("没有可写入的模板数据")
    try:
        workbook = load_workbook(source)
    except OSError as exc:
        raise GaiaTemplateWriteError("无法打开盖亚模板") from exc
    sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
    header = _header_row(sheet)
    first_data_row = header + 1
    _assert_empty_data_area(sheet, first_data_row)
    reference_row = first_data_row
    for offset, row in enumerate(materialized):
        row_number = first_data_row + offset
        for column, value in enumerate(row.values(), start=1):
            target = sheet.cell(row_number, column, value=value)
            if reference_row <= sheet.max_row:
                source_cell = sheet.cell(reference_row, column)
                if source_cell.has_style:
                    target._style = copy(source_cell._style)
                if source_cell.number_format:
                    target.number_format = source_cell.number_format
                if source_cell.alignment:
                    target.alignment = copy(source_cell.alignment)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return str(destination)
